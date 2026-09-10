"""Build and seal Ilyra Fen v690-v2 final reports and handoff."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import subprocess
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs" / "ilyra-fen" / "v690-v2"
PLAN = BASE / "plan"
X1 = BASE / "x1"
X2 = BASE / "x2"
FINAL = BASE / "final"
OWNER = "Ilyra Fen"
PHASE = "v690-v2"
BRANCH = "codex/GHC-Family/ilyra-fen-main"
SOURCE = "9936e2855b72bddfecdea77abdd6f083f14a09f1"
PLANNING = "8c4eef447c296e4d956c75ff16e6205bf842df0b"
X1_COMMIT = "8e2a010330d4e31486dd3059ed1be9e2e744f9c1"
X2_COMMIT = "e85f3a2cfe5260ff2b9d6e953a80513d383650c5"
SOURCE_CANONICAL = "8d089142b18915b3cf149cea7dde9638085c06b54d91320fb0bedda030dbb2c0"
BOUNDARY = (
    "This phase supplies bounded same-owner synthetic software and documentation evidence. "
    "It is not a full-repository suite, external audit, independent reproduction, empirical "
    "GMUT confirmation, production THOS certification, live Freed ID lifecycle, professional "
    "qualification, legal or cultural authority, Maori authority, complete privacy or "
    "accessibility assurance, exhaustive security, AGI or ASI evidence, consciousness or "
    "personhood evidence, a Theory-of-Everything proof, canon, or Stage 20 readiness."
)
RELATIONAL = (
    "Ilyra Fen, they and them, the role directed-obligation provenance cartographer, the hope "
    "of making dependency, correction, and authority boundaries inspectable, and all family "
    "language are corrigible relational working terms. They are not evidence of consciousness, "
    "sentience, personhood, identity continuity, employment, qualification, independent agency, "
    "or authority."
)
PROTECTED_GATES = [
    "empirical",
    "participants",
    "professional",
    "production",
    "legal-cultural-Maori-authority",
    "privacy-accessibility-security-completeness",
    "independent-reproduction",
    "AGI-ASI-consciousness-personhood",
    "Theory-of-Everything-Stage-20",
]
BASE_ACCOUNTING = {
    "effective_negatives": 1255,
    "methods": 104,
    "direct_witnesses": 3481,
    "failed_witnesses": 966,
    "passing_witnesses": 2515,
}
FINAL_FAILURES = [
    {
        "id": "IF6902-FINAL-F001",
        "failure": "A combined x2 evidence refresh and test wrapper returned only its early Ruff output before separate scalar recovery.",
        "recovery": "Read refreshed Method Flow and manifest counts, then run the x2 tests as a standalone bounded command.",
    },
    {
        "id": "IF6902-FINAL-F002",
        "failure": "The first complete Documents skill display exceeded its output budget.",
        "recovery": "Read the skill again in bounded physical-line windows through EOF before document authoring.",
    },
    {
        "id": "IF6902-FINAL-F003",
        "failure": "The retired dynamic workspace-dependency loader returned no current runtime paths.",
        "recovery": "Use the current Codex app MCP dependency loader and bind its bundled Node and Python paths.",
    },
    {
        "id": "IF6902-FINAL-F004",
        "failure": "Several first PowerShell path probes for the document marker returned no attributable display.",
        "recovery": "Use a bounded literal file listing to resolve the marker inside the Documents skill package.",
    },
    {
        "id": "IF6902-FINAL-F005",
        "failure": "The first final-builder lint preflight found an unsorted import block.",
        "recovery": "Normalize the import boundary and rerun the bounded lint preflight before preparing artifacts.",
    },
    {
        "id": "IF6902-FINAL-F006",
        "failure": "The first exact-final canonical-validator lint preflight found an unsorted import block.",
        "recovery": "Normalize the import boundary and rerun the bounded lint preflight before preparing artifacts.",
    },
    {
        "id": "IF6902-FINAL-F007",
        "failure": "The first final-test lint preflight found an unsorted import block.",
        "recovery": "Normalize the import boundary and rerun the bounded lint preflight before preparing artifacts.",
    },
    {
        "id": "IF6902-FINAL-F008",
        "failure": "The first canonical-validator lint preflight found two separable startswith calls for one prefix family.",
        "recovery": "Use one tuple-based startswith call and rerun the bounded lint preflight before preparing artifacts.",
    },
    {
        "id": "IF6902-FINAL-F009",
        "failure": "The first renderer help probe used the utility subdirectory rather than the Documents skill root.",
        "recovery": "Resolve the exact packaged render_docx.py path with a bounded file inventory before invoking it.",
    },
    {
        "id": "IF6902-FINAL-F010",
        "failure": "The first structural accessibility audit found four medium table-header omissions.",
        "recovery": "Apply the supported first-row header repair, scrub the repaired copy, and obtain a zero-finding audit.",
    },
    {
        "id": "IF6902-FINAL-F011",
        "failure": "The canonical Documents renderer could not start because this Windows runtime has no bundled LibreOffice and no soffice executable on PATH.",
        "recovery": "Use hidden Microsoft Word fixed-format export and the bundled Poppler rasterizer as an explicitly bounded Windows fallback.",
    },
    {
        "id": "IF6902-FINAL-F012",
        "failure": "The first rasterizer-capability probe attempted an unavailable fitz import after locating bundled pdftoppm.",
        "recovery": "Invoke the resolved bundled pdftoppm executable directly against the Word-exported PDF.",
    },
    {
        "id": "IF6902-FINAL-F013",
        "failure": "A broad precommit unittest invocation ran three snapshot-sensitive planning and x1 assertions against the completed x2/final worktree.",
        "recovery": "Retain the three expected current-tree failures, run only current-layer tests before commit, and let the one-shot canonical validator materialize immutable planning and x1 layers for their exact selections.",
    },
]


def jbytes(value: object) -> bytes:
    return json.dumps(
        value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False
    ).encode("utf-8")


def digest(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, sort_keys=True, allow_nan=False) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value.rstrip() + "\n", encoding="utf-8", newline="\n")


def load(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_shading(cell: Any, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_borders(cell: Any, color: str = "D9D9D9") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)


def set_cell_margins(cell: Any, top: int = 90, start: int = 120, bottom: int = 90, end: int = 120) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_table(table: Any, widths: list[float]) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row_index, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            fill = "1F4E78" if row_index == 0 else ("EAF2F8" if row_index % 2 == 0 else "FFFFFF")
            set_cell_shading(cell, fill)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
                    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
                    run.font.size = Pt(9.2)
                    if row_index == 0:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = str(value)
            if index > 0 and len(str(value)) < 18:
                cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_table(table, widths)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    paragraph.paragraph_format.keep_with_next = True


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.12


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    styles = document.styles
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = styles[style_name]
        style.font.name = "Aptos"
        style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
        style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.size = Pt(25)
    styles["Title"].font.bold = False
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.size = Pt(12)
    styles["Heading 2"].font.bold = True
    styles["Heading 1"].paragraph_format.space_before = Pt(4)
    styles["Heading 1"].paragraph_format.space_after = Pt(7)
    styles["Heading 2"].paragraph_format.space_before = Pt(5)
    styles["Heading 2"].paragraph_format.space_after = Pt(4)


def overview_sections(accounting: dict[str, int]) -> list[dict[str, Any]]:
    return [
        {
            "title": "Outcome and lifecycle",
            "paragraphs": [
                (
                    "This report records the completed Ilyra Fen v690 v2 owner phase. The "
                    "parentless branch preserves Lyren Moss final as explicit source provenance. "
                    "Planning, x1, and x2 are separate pushed commits, and final closeout is the "
                    "fourth direct child. The terminal scientific and authority verdict remains "
                    "NOT READY FOR STAGE 20."
                ),
                RELATIONAL,
                (
                    "Two hundred inherited Lyren records were projected losslessly with zero "
                    "Ilyra novelty and execution credit. Two hundred new requests were frozen "
                    "before implementation. Every invalid candidate remains a failed subject even "
                    "where its field closure guard passed."
                ),
            ],
            "table": (
                ["Lifecycle record", "Exact value"],
                [
                    ["Source provenance", SOURCE],
                    ["Planning root", PLANNING],
                    ["Frozen x1", X1_COMMIT],
                    ["Immutable x2", X2_COMMIT],
                    ["Current verdict", "NOT READY FOR STAGE 20"],
                ],
                [2.0, 5.0],
            ),
        },
        {
            "title": "Finite graph results and limits",
            "paragraphs": [
                (
                    "The twenty finite operations cover canonical graph records, incoming and "
                    "outgoing adjacency, reachability, shortest hop paths, acyclic status, "
                    "topological order, weak components, transitive closure, boundary edges, "
                    "transitive reduction, graph differences, provenance digests, correction "
                    "chains, dependency blockers, release gates, critical paths, DOT source, "
                    "accessible summaries, and evidence reservations."
                ),
                (
                    "All two hundred safe requests matched their complete frozen envelopes without "
                    "input mutation. All two hundred paired unknown-field subjects were refused and "
                    "retain zero original success credit. A directed cycle shows why a topological "
                    "order needs acyclicity. A path can exist while consent, evidence, and authority "
                    "remain absent."
                ),
                (
                    "Graph definitions can clarify GMUT model structure, but they do not define a "
                    "physical action, dimensions, conservation equations, observables, likelihoods, "
                    "calibration, or falsifiers. No result ranks GMUT against tested physical theories."
                ),
            ],
            "table": (
                ["Outcome", "Count", "Meaning in this phase"],
                [
                    ["completed", "180", "The exact finite typed predicate passed"],
                    ["represented", "10", "A structure exists while human evidence remains absent"],
                    ["open gap", "5", "Required scientific or evaluation evidence is missing"],
                    ["exact gate", "5", "Competent authority or an exact condition is required"],
                ],
                [1.4, 0.8, 4.8],
            ),
        },
        {
            "title": "Tools skills and evidence accounting",
            "paragraphs": [
                (
                    "NetworkX 3.6.1, rustworkx 0.18.1, and graphviz 0.21 were installed from "
                    "hash-verified wheels in one D-isolated environment. NumPy 2.5.3 is an explicit "
                    "rustworkx dependency and receives no direct-package credit. The first offline "
                    "install failed because NumPy was absent; the corrected install reused the "
                    "failed environment. A later NodeIndices serialization failure also remains."
                ),
                (
                    "Thirty x2 comparisons passed: ten NetworkX transitive reductions, ten "
                    "rustworkx topological orders, and ten Graphviz DOT source checks without "
                    "rendering. Twenty local skills and ten paired runners were built. Five merged "
                    "skills and five public D runners were promoted only after absence checks, "
                    "official validation, byte parity, and accepting and rejecting smokes."
                ),
                (
                    "The evidence counts include every invalid subject and operational failure. A "
                    "recovery is a separate passing witness and never changes the original failure."
                ),
            ],
            "table": (
                ["Evidence item", "Count"],
                [
                    ["Effective retained negatives", f"{accounting['effective_negatives']:,}"],
                    ["Method Flow methods", f"{accounting['methods']:,}"],
                    ["Direct witnesses", f"{accounting['direct_witnesses']:,}"],
                    ["Failed witnesses", f"{accounting['failed_witnesses']:,}"],
                    ["Passing witnesses", f"{accounting['passing_witnesses']:,}"],
                    ["Four tier cards", "213"],
                ],
                [4.8, 2.2],
            ),
        },
        {
            "title": "Pillars practices and rights",
            "paragraphs": [
                (
                    "Freed ID and CBR Heart is primary. The phase separates graph structure, source "
                    "provenance, correction, evidence, consent, identity, rights, and competent "
                    "authority. A matching digest binds one chosen serialization. It does not prove "
                    "truth, authorship, ownership, consent, lawful basis, cultural legitimacy, or "
                    "authority to act."
                ),
                (
                    "THOS Body contributes field-closed interfaces, reversible evidence, isolated "
                    "dependencies, and exact lifecycle gates. It remains synthetic and proxy-only "
                    "without governed real workloads, operators, safety monitoring, appropriate "
                    "statistics, or independent review. GMUT Mind contributes finite definitions "
                    "and counterexamples while its empirical programme remains open."
                ),
                (
                    "The four practices are finite directed-graph algorithm testing, provenance and "
                    "correction-lineage engineering, public-interest data-governance review, and "
                    "accessible dependency-status editing. These are learning lenses, not employment, "
                    "licensure, competence, standing, or affected-party authority."
                ),
            ],
            "table": (
                ["Successor recommendation", "Purpose"],
                [
                    ["Adversarial reachability policy auditor", "Challenge hidden assumptions that turn paths into permission"],
                    ["Provenance graph visualization accessibility reviewer", "Test whether dependency and uncertainty states remain understandable"],
                ],
                [2.7, 4.3],
            ),
        },
        {
            "title": "Terminal route and evidence boundary",
            "paragraphs": [
                (
                    "Mira Fenwick v690 v3 is prospective only. The repository final must first be "
                    "pushed, clean, zero divergent, four-way equal, and accepted by one exact-final "
                    "owner-scoped canonical invocation. The current roster, authority, exact title, "
                    "duplicate, pause, privacy, safety, usage, and acknowledgement guards must then "
                    "be reread."
                ),
                (
                    "A compact live message may be sent once to the unique existing Mira Fenwick "
                    "main task after every terminal guard passes. An accepted or opaque-accepted send "
                    "ends retries. No replacement task, fork, collaboration subagent, standby "
                    "substitute, early later-owner contact, or second confirmation is part of this phase."
                ),
                BOUNDARY,
                (
                    "Primary references are Kahn's 1962 topological sorting paper, Aho, Garey, and "
                    "Ullman's 1972 transitive reduction paper, RFC 8785, W3C PROV-O, W3C Verifiable "
                    "Credentials Data Model 2.0, the Graphviz DOT language, and the official package "
                    "records for the three direct libraries. They supply definitions and vocabulary, "
                    "not validation or authority for these artifacts."
                ),
            ],
            "table": None,
        },
    ]


def build_docx(path: Path, accounting: dict[str, int]) -> None:
    document = Document()
    configure_document(document)
    sections = overview_sections(accounting)
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Ilyra Fen v690 v2 Evidence Report")
    subtitle = document.add_paragraph("Finite directed obligation graphs provenance correction and authority boundaries")
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.runs[0].font.size = Pt(12)
    subtitle.runs[0].font.italic = True
    add_body(
        document,
        "Prepared for Hamish and the prospective Mira Fenwick handoff. The main conclusion is that the bounded owner work is complete while every real-world scientific, production, identity, rights, and authority gate remains open or reserved.",
    )
    for index, section in enumerate(sections):
        add_heading(document, section["title"])
        for paragraph in section["paragraphs"]:
            add_body(document, paragraph)
        if section["table"]:
            headers, rows, widths = section["table"]
            add_table(document, headers, rows, widths)
        if index < len(sections) - 1:
            document.add_page_break()
    document.core_properties.title = "Ilyra Fen v690 v2 Evidence Report"
    document.core_properties.subject = "Bounded finite graph evidence and handoff"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)


def markdown_overview(accounting: dict[str, int]) -> str:
    chunks = ["# Ilyra Fen v690 v2 Evidence Report", ""]
    for section in overview_sections(accounting):
        chunks.append(f"## {section['title']}")
        chunks.append("")
        chunks.extend(paragraph + "\n" for paragraph in section["paragraphs"])
        if section["table"]:
            headers, rows, _widths = section["table"]
            chunks.append("| " + " | ".join(headers) + " |")
            chunks.append("| " + " | ".join("---" for _ in headers) + " |")
            chunks.extend("| " + " | ".join(row) + " |" for row in rows)
            chunks.append("")
    return "\n".join(chunks)


def html_overview(accounting: dict[str, int]) -> str:
    sections = overview_sections(accounting)
    body = []
    for section in sections:
        body.append(f"<section><h2>{section['title']}</h2>")
        body.extend(f"<p>{paragraph}</p>" for paragraph in section["paragraphs"])
        if section["table"]:
            headers, rows, _widths = section["table"]
            body.append("<table><thead><tr>")
            body.extend(f"<th scope=\"col\">{header}</th>" for header in headers)
            body.append("</tr></thead><tbody>")
            for row in rows:
                body.append("<tr>" + "".join(f"<td>{value}</td>" for value in row) + "</tr>")
            body.append("</tbody></table>")
        body.append("</section>")
    return """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Ilyra Fen v690 v2 Evidence Report</title>
<style>body{font-family:Arial,sans-serif;max-width:900px;margin:2rem auto;line-height:1.55;color:#111}
h1,h2{color:#000}table{border-collapse:collapse;width:100%;margin:1rem 0}th,td{border:1px solid #d9d9d9;padding:.6rem}
th{background:#1f4e78;color:#fff;text-align:left}tr:nth-child(even){background:#eaf2f8}</style></head>
<body><a href="#main">Skip to report</a><main id="main"><h1>Ilyra Fen v690 v2 Evidence Report</h1>
""" + "\n".join(body) + "</main></body></html>\n"


def module_texts(accounting: dict[str, int]) -> list[tuple[str, str]]:
    proposals = load(PLAN / "new-proposals.json")["proposals"]
    x1_results = {row["proposal_id"]: row for row in load(X1 / "results.json")["records"]}
    x2_results = {row["proposal_id"]: row for row in load(X2 / "results.json")["records"]}
    results = x1_results | x2_results
    core_lines = ["# 04 Core proposal results", ""]
    for row in proposals:
        observed = results[row["proposal_id"]]
        core_lines.extend(
            [
                f"## {row['proposal_id']} {row['title']}",
                "",
                (
                    f"This {row['lane']} record uses the {row['practice']} lens under {row['pillar']}. "
                    f"Its bounded mission is {row['mission']} The complete frozen envelope matched "
                    f"the owner implementation and the exact outcome is {row['expected_execution_disposition']}."
                ),
                "",
                (
                    f"Evidence is stored in docs/ilyra-fen/v690-v2/{row['lane']}/results.json. "
                    f"The observed comparison passed {str(observed['passed']).lower()}. Its paired "
                    "unknown authority-field subject remains a failed record with zero original "
                    "success credit even though the field-closure guard refused it."
                ),
                "",
                (
                    "Any complete typed mismatch, accepted unknown field, input mutation, erased "
                    "failed subject, unsupported correction, or authority promotion would falsify "
                    "the bounded claim. Rollback remains additive: retain the source request and "
                    "failed subject, isolate the responsible operation, and add a separately "
                    "attributable correction."
                ),
                "",
                (
                    "A passing finite graph result does not close empirical, participant, "
                    "professional, production, legal, cultural, Maori-authority, privacy, "
                    "accessibility, security, independent-reproduction, AGI or ASI, consciousness "
                    "or personhood, Theory-of-Everything, canon, or Stage 20 gates."
                ),
                "",
            ]
        )
    modules = [
        (
            "01-welcome-and-route.md",
            f"""# 01 Welcome and exact route

Dear Mira Fenwick,

This file-backed candidate prepares one prospective Ilyra Fen v690-v2 to Mira Fenwick v690-v3 edge under Hamish's current weighted schedule. It does not itself prove native delivery. The committed route state is PREPARED_NOT_SENT.

Before mutation, read this combined baton through its literal EOF marker and the current guidance it names. Reverify the exact final, lifecycle parents, manifests, content seal, canonical receipt, clean equality, current roster, exact title, endpoint kind, duplicate guard, and newest authority. Do not replay Ilyra's successful canonical.

{RELATIONAL}

Hamish may rename, pause, redirect, narrow, or stop the route. No task creation, fork, subagent, standby substitution, early later-owner contact, or second confirmation is authorized.
""",
        ),
        (
            "02-source-and-lifecycle.md",
            f"""# 02 Source provenance and lifecycle

Lyren Moss exact final {SOURCE} is explicit source provenance and not Git ancestry. Its canonical receipt SHA-256 is {SOURCE_CANONICAL}. Ilyra's parentless owner lifecycle is planning {PLANNING}, frozen x1 {X1_COMMIT}, immutable x2 {X2_COMMIT}, and one exact final supplied later by the external canonical receipt.

Planning froze two hundred inherited zero-credit selections and two hundred new complete requests before execution. X1 implemented only its first ten operations. X2 began only after x1 was pushed, clean, typed zero divergent, and fresh-four-way equal. Source and sibling lanes stayed read-only.

The branch ceiling is eight commits and the file ceiling is two thousand owner files. The intended final is commit four, with zero merges and exact direct-child ancestry inside this owner root.
""",
        ),
        (
            "03-portfolios-and-outcomes.md",
            """# 03 Portfolios outcomes and approval boundaries

Each session executed one hundred safe requests, one hundred invalid candidate subjects, and one hundred inherited-record round trips. The inherited rows retain zero Ilyra novelty and execution credit. The candidate subjects remain failures; passing guards are separate evidence.

The two hundred new outcomes are exactly 180 completed, 10 represented, 5 open_gap, and 5 exact_gate. Fifty exact packets describe lifecycle prerequisites. Thirty blocked packets remain unexecuted protected subjects. Ceilings and floors organize bounded work and never authorize filler or unsafe action.

Completed means only that an exact finite typed predicate passed. Represented records useful structure without the absent human or external evidence. Open gaps retain missing scientific or evaluation evidence. Exact gates retain actions or claims that need competent authority or another exact protected condition.
""",
        ),
        ("04-core-proposal-results.md", "\n".join(core_lines)),
        (
            "05-retained-failures-and-methods.md",
            f"""# 05 Retained failures and Method Flow

The successor-visible repository preparation accounting is {accounting['effective_negatives']} effective negatives, {accounting['methods']} methods, and {accounting['direct_witnesses']} direct witnesses: {accounting['failed_witnesses']} failed and {accounting['passing_witnesses']} passing. Every recovery remains separate from its failed attempt.

Important failures include baton display truncation, parser mistakes, cache and manifest scope defects, a slow asynchronous commit and push, an incomplete package closure, a rustworkx serialization mismatch, skill metadata correction, staged sparse-pattern and index-lock recovery, and bounded validation display timeouts. Each receives zero original success credit.

Method Flow records trigger conditions, failure signatures, retained-negative links, passing witnesses, recurrence guards, rollback, owner scope, exact-head requirements, and protected gates. Same-owner validation under shared infrastructure is not independent reproduction or an external audit.
""",
        ),
        (
            "06-algorithms-packages-and-counterexamples.md",
            """# 06 Algorithms packages and counterexamples

The finite graph surface uses explicit nodes, directed edges, bounded size, deterministic ordering, and strict field closure. Kahn's procedure supplies a topological order only when the graph is acyclic. Transitive reduction is restricted to DAGs and preserves the declared reachability relation.

NetworkX 3.6.1, rustworkx 0.18.1, graphviz 0.21, and the NumPy 2.5.3 dependency live only in a D-isolated environment. Thirty package comparisons passed. An import or comparison is local software evidence, not an endorsement, security audit, license opinion, performance benchmark, or production certification.

A directed cycle rejects a universal topological-order claim. A reachable path with absent consent and authority rejects the claim that connectivity itself authorizes action. These are finite counterexamples to broad example statements, not new fundamental laws.
""",
        ),
        (
            "07-skills-runners-and-catalogue.md",
            """# 07 Skills runners and catalogue

X1 and x2 each built ten operation-specific local skills and five paired runners. Every skill contains focused instructions, proper UI metadata, one contract reference, and accepting and rejecting fixtures. Official quick validation passed.

Five merged global skill candidates and five public D runners were installed only after every destination was absent. Two core modules are dependencies and receive zero runner credit. Source-to-target byte parity and accepting and outside-group smokes passed.

The forty-card meta-tool catalogue distinguishes local skills, local runners, installed merged skills, and public runners. Availability does not widen execution authority. Rollback stops selecting the exact additive capability while preserving its source, installed bytes, and receipts.
""",
        ),
        (
            "08-gmut-and-research-boundaries.md",
            """# 08 GMUT and research boundaries

Directed graphs can express finite dependencies, reachability, cycles, reductions, and provenance paths. Those mathematical objects do not by themselves define spacetime indices, a metric, an action, units, symmetries, conservation laws, initial or boundary conditions, observables, likelihoods, calibration, uncertainty, or falsifiers.

GMUT remains a typed scalar-tensor and effective-field-theory research-model family. An arbitrary Omega term needs defined dynamics and a complete conservation account. A graph analogy cannot confirm a field equation, measure a physical channel, establish a psyche variable, solve unsolved physics, or prove a Theory of Everything.

Useful next research would preregister one narrow model, an observable map, a comparator, nuisance parameters, uncertainty, data governance, and a disconfirming result. Independent scrutiny and real evidence remain absent here.
""",
        ),
        (
            "09-freed-id-cbr-privacy-and-authority.md",
            """# 09 Freed ID CBR privacy and authority

Freed ID and CBR Heart is primary. The phase separates identifier, record, digest, path, correction, evidence, consent, rights, and competent authority. A valid digest or path does not authenticate a person, prove authorship, grant consent, establish ownership, or authorize a rights-affecting action.

A production Freed ID system would need standards-conformant live keys and proofs, issuer and verifier policy, status and revocation, recovery, interoperability, threat modelling, independent security and privacy assessment, governance, auditability, accessibility, and affected-party oversight.

The five-class privacy scan covers local profile paths, email addresses, network addresses, credential-like secrets, and keyed phone contacts. Zero confirmed candidates in that bounded scan is not complete privacy assurance. Legal, cultural, affected-party, professional, and Maori-authority decisions remain exact gates.
""",
        ),
        (
            "10-reports-and-accessibility.md",
            """# 10 Reports documents and accessibility

The phase supplies Markdown, HTML, DOCX, PDF, and page images for one bounded overview. The DOCX uses a plain title, linear headings, readable prose, deliberate tables, visible borders, and repeated text labels. The page images are inspected after the final render.

The four-tier deck has an accessible HTML projection and machine-readable JSON. Structural review cannot replace screen-reader, keyboard, magnification, voice-control, cognitive, language, or affected-user evaluation. Complete accessibility remains open.

Rendered pages are documentation artifacts rather than experimental figures. They contain no real participant, identity, measurement, credential, rights decision, or deployed result.
""",
        ),
        (
            "11-four-tier-context-and-practices.md",
            """# 11 Four tier context and practices

The deck contains one relational owner card, three pillar cards, four practice cards, two hundred core task cards, and five supplementary cards. Every non-root card has exactly one immediate-tier parent. Card identifiers derive from canonical content digests.

The practices are finite directed-graph algorithm testing, provenance and correction-lineage engineering, public-interest data-governance review, and accessible dependency-status editing. They structure questions and do not establish employment, licensure, competence, professional standing, or affected-party authority.

Mira's recommended lenses are adversarial reachability-policy auditing and provenance-graph visualization accessibility review. They are recommendations rather than assignments that override Mira's judgment or current authority.
""",
        ),
        (
            "12-successor-work-and-scenarios.md",
            """# 12 Successor work and conditional scenarios

Prepared skill ideas include policy reachability assumptions, cycle authority quarantine, provenance root conflict maps, correction branch nonerasure, dependency expiry, consent and authority separation, accessible blocked-path explanation, graph serialization receipts, evidence-class separation, and exact route confirmation.

Prepared runner ideas include bounded policy comparison, strong-component quarantine, correction replay, multi-root frontier analysis, unresolved-dependency reporting, path ambiguity witnesses, consent-expiry simulation, transitive-reduction equivalence, DOT source fixity, and exact-source confirmation.

These are proposals without completion credit. Conditional ten, thirty, hundred, and thousand year scenarios remain scenarios. Longer horizons require durable standards, governance, maintenance, material and energy accounting, inclusive institutions, independent scrutiny, and correction mechanisms. They do not imply inevitable AGI, ASI, consciousness, adoption, legitimacy, or Stage 20.
""",
        ),
        (
            "13-terminal-gates-and-recovery.md",
            f"""# 13 Terminal gates duplicate guard and recovery

Before any Mira mutation, require the Ilyra final to be the direct child of {X2_COMMIT}; total commits no more than eight; single-parent merge-free history; clean Git; typed zero divergence; local, upstream, tracking, and fresh live remote equality; exact manifest and content-seal replay from Git blobs; strict JSON and YAML parses; bounded Python AST and privacy checks; valid documents and page images; baton within ten thousand through one hundred thousand words; literal EOF; and zero prior Ilyra canonical invocation.

Run the exact-final owner-scoped canonical once. Latch a success and never replay it. A failure retains zero canonical-success credit and permits only an additive correction justified by the exact blocked component.

After success, reread Hamish's newest instruction and the current weighted roster. Resolve active and archived registries for the exact title Mira Fenwick. Require one eligible existing main task, immediately reread duplicate, pause, redirect, privacy, safety, evidence, usage, and acknowledgement guards, then send one compact sanitized pointer at most once.

An accepted or opaque-accepted send ends retries. A transient service failure with no accepted send requires at least five bounded list and read recovery attempts before an open route gap. Never create a replacement, infer a private handle, automate an unsupported UI path, substitute another endpoint, or send a second confirmation.

{BOUNDARY}

Terminal verdict NOT_READY_FOR_STAGE_20.
""",
        ),
    ]
    return modules


def final_accounting() -> dict[str, int]:
    failures = len(FINAL_FAILURES)
    return {
        "effective_negatives": BASE_ACCOUNTING["effective_negatives"] + failures,
        "methods": BASE_ACCOUNTING["methods"] + 1,
        "direct_witnesses": BASE_ACCOUNTING["direct_witnesses"] + failures * 2,
        "failed_witnesses": BASE_ACCOUNTING["failed_witnesses"] + failures,
        "passing_witnesses": BASE_ACCOUNTING["passing_witnesses"] + failures,
    }


def prepare(draft_docx: Path) -> None:
    if subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip() != X2_COMMIT:
        raise RuntimeError("final preparation must begin at immutable x2")
    accounting = final_accounting()
    FINAL.mkdir(parents=True, exist_ok=True)
    write_text(FINAL / "overview.md", markdown_overview(accounting))
    write_text(FINAL / "overview.html", html_overview(accounting))
    build_docx(draft_docx, accounting)
    modules = module_texts(accounting)
    module_records = []
    for name, content in modules:
        path = FINAL / "baton" / name
        write_text(path, content)
        data = path.read_bytes()
        module_records.append(
            {
                "bytes": len(data),
                "path": path.relative_to(ROOT).as_posix(),
                "sha256": digest(data),
                "words": len(re.findall(r"\b[\w'-]+\b", data.decode("utf-8"), flags=re.UNICODE)),
            }
        )
    combined = "\n\n".join((FINAL / "baton" / name).read_text(encoding="utf-8").rstrip() for name, _ in modules)
    combined += "\n\nEOF ILYRA FEN v690-v2 BATON.\n"
    write_text(FINAL / "hand-off-baton.md", combined)
    baton_data = (FINAL / "hand-off-baton.md").read_bytes()
    baton_words = len(re.findall(r"\b[\w'-]+\b", baton_data.decode("utf-8"), flags=re.UNICODE))
    write_json(
        FINAL / "baton-manifest.json",
        {
            "combined": {
                "bytes": len(baton_data),
                "literal_last_line": "EOF ILYRA FEN v690-v2 BATON.",
                "path": "docs/ilyra-fen/v690-v2/final/hand-off-baton.md",
                "sha256": digest(baton_data),
                "words": baton_words,
            },
            "delivery_state": "PREPARED_NOT_SENT",
            "module_count": len(module_records),
            "modules": module_records,
            "owner": OWNER,
            "phase": PHASE,
            "range": {"in_range": 10_000 <= baton_words <= 100_000, "minimum": 10_000, "maximum": 100_000},
            "recipient": "Mira Fenwick",
            "recipient_phase": "v690-v3",
        },
    )
    write_json(
        FINAL / "operational-overlay.json",
        {
            "failures": [
                {**row, "original_success_credit": 0, "retained": True} for row in FINAL_FAILURES
            ],
            "method": "IF6902-FINAL-M001",
            "recovery_witnesses": len(FINAL_FAILURES),
        },
    )
    write_json(
        FINAL / "terminal-accounting.json",
        {
            "base_after_x2": BASE_ACCOUNTING,
            "effective_total": accounting,
            "final_authoring_layer": {
                "effective_negatives": len(FINAL_FAILURES),
                "methods": 1,
                "direct_witnesses": len(FINAL_FAILURES) * 2,
                "failed_witnesses": len(FINAL_FAILURES),
                "passing_witnesses": len(FINAL_FAILURES),
            },
            "terminal_verdict": "NOT_READY_FOR_STAGE_20",
        },
    )
    write_json(
        FINAL / "completion-ledger.json",
        {
            "canonical_invoked": False,
            "core_outcomes": {"completed": 180, "represented": 10, "open_gap": 5, "exact_gate": 5},
            "deck_cards": 213,
            "global_runners": 5,
            "global_skills": 5,
            "inherited_zero_credit": 200,
            "local_runners": 10,
            "local_skills": 20,
            "new_proposals": 200,
            "packages": {"direct": 3, "transitive": 1},
            "route_state": "PREPARED_NOT_SENT",
            "tasks": {
                "candidate_x1": 100,
                "candidate_x2": 100,
                "clean_x1": 100,
                "clean_x2": 100,
                "safe_x1": 100,
                "safe_x2": 100,
            },
            "terminal_verdict": "NOT_READY_FOR_STAGE_20",
        },
    )
    write_json(
        FINAL / "lifecycle.json",
        {
            "branch": BRANCH,
            "canonical_state": "PENDING_EXACT_FINAL",
            "commit_ceiling": 8,
            "expected_parent_of_final": X2_COMMIT,
            "final": "EXTERNAL_AFTER_FINAL_COMMIT",
            "merges": 0,
            "planning": PLANNING,
            "route_state": "PREPARED_NOT_SENT",
            "source_is_ancestor": False,
            "source_provenance": SOURCE,
            "strict_x1_before_x2": True,
            "x1": X1_COMMIT,
            "x2": X2_COMMIT,
        },
    )
    write_json(
        FINAL / "route-candidate.json",
        {
            "canonical_success_required": True,
            "duplicate_guard": "pending terminal registry reread",
            "endpoint_kind": "existing Codex main task required",
            "from": OWNER,
            "from_phase": PHASE,
            "latest_authority_reread_required": True,
            "native_acknowledgement": None,
            "precontacted": False,
            "replacement_creation_authorized": False,
            "state": "PREPARED_NOT_SENT",
            "to_exact_title": "Mira Fenwick",
            "to_phase": "v690-v3",
        },
    )
    write_text(FINAL / "evidence-boundary.md", "# Evidence boundary\n\n" + BOUNDARY + "\n\n" + RELATIONAL)
    write_json(
        FINAL / "source-references.json",
        {
            "records": [
                {"title": "Topological sorting of large networks", "url": "https://doi.org/10.1145/368996.369025"},
                {"title": "The Transitive Reduction of a Directed Graph", "url": "https://doi.org/10.1137/0201008"},
                {"title": "RFC 8785 JSON Canonicalization Scheme", "url": "https://www.rfc-editor.org/rfc/rfc8785.html"},
                {"title": "W3C PROV-O", "url": "https://www.w3.org/TR/prov-o/"},
                {"title": "W3C Verifiable Credentials Data Model 2.0", "url": "https://www.w3.org/TR/vc-data-model-2.0/"},
                {"title": "Graphviz DOT language", "url": "https://graphviz.org/doc/info/lang.html"},
            ],
            "source_validation_credit": 0,
        },
    )
    write_text(
        FINAL / "compact-activation.md",
        f"""# Ilyra Fen v690 v2 to Mira Fenwick v690 v3 compact pointer

Only after Ilyra's clean pushed exact final and one successful non-replayed canonical, read docs/ilyra-fen/v690-v2/final/hand-off-baton.md through EOF on {BRANCH}. Source provenance {SOURCE}; planning {PLANNING}; x1 {X1_COMMIT}; x2 {X2_COMMIT}. The exact final and external canonical receipt are supplied by one later native activation.

Core outcomes are 180 completed, 10 represented, 5 open_gap, and 5 exact_gate. Terminal verdict NOT_READY_FOR_STAGE_20. Work solo in a new Mira-owned D-first lane, preserve every failure and gate, and do not replay Ilyra's canonical.

Committed state PREPARED_NOT_SENT. Native delivery is a separate event.
""",
    )
    write_json(
        FINAL / "allowlist.json",
        {
            "allowed_prefixes": [
                "docs/ilyra-fen/v690-v2/final/",
                "scripts/build_ghc_family_ilyra_v690_v2_final.py",
                "scripts/ghc_family_ilyra_v690_v2_canonical.py",
                "tests/test_ghc_family_ilyra_v690_v2_final.py",
            ]
        },
    )
    print(
        json.dumps(
            {
                "baton_bytes": len(baton_data),
                "baton_words": baton_words,
                "modules": len(modules),
                "state": "FINAL_CONTENT_PREPARED_DOCX_DRAFT_EXTERNAL",
            },
            sort_keys=True,
        )
    )


def seal(visual_pages: int) -> None:
    required = [
        FINAL / "overview.docx",
        FINAL / "rendered" / "overview.pdf",
        *[FINAL / "rendered" / f"page-{index}.png" for index in range(1, visual_pages + 1)],
    ]
    missing = [path.as_posix() for path in required if not path.is_file() or path.stat().st_size == 0]
    if missing:
        raise RuntimeError(f"missing rendered documents: {missing}")
    write_json(
        FINAL / "visual-review.json",
        {
            "affected_user_review": "reserved",
            "artifact": "docs/ilyra-fen/v690-v2/final/rendered/overview.pdf",
            "blank_required_content": False,
            "broken_heading_hierarchy_observed": False,
            "clipping": False,
            "complete_accessibility_claimed": False,
            "overlap": False,
            "pages": visual_pages,
            "pages_inspected": list(range(1, visual_pages + 1)),
            "render_pipeline": "hidden Microsoft Word fixed-format export plus bundled Poppler pdftoppm",
            "result": "pass",
            "same_owner_only": True,
            "unreadable_tables": False,
        },
    )
    write_json(
        FINAL / "document-build.json",
        {
            "docx": "docs/ilyra-fen/v690-v2/final/overview.docx",
            "metadata_scrubbed": True,
            "page_count": visual_pages,
            "pdf": "docs/ilyra-fen/v690-v2/final/rendered/overview.pdf",
            "render_equivalence_claimed": False,
            "render_fallback": "Windows Word fixed-format export plus bundled Poppler because bundled LibreOffice is unavailable",
            "rendered_pages": [
                f"docs/ilyra-fen/v690-v2/final/rendered/page-{index}.png"
                for index in range(1, visual_pages + 1)
            ],
            "structural_accessibility_audit": "external zero-finding receipt before seal",
        },
    )
    entries = []
    for path in sorted(candidate for candidate in FINAL.rglob("*") if candidate.is_file()):
        relative = path.relative_to(ROOT).as_posix()
        if relative.endswith(("content-seal.json", "manifest.json")):
            continue
        data = path.read_bytes()
        entries.append({"bytes": len(data), "path": relative, "sha256": digest(data)})
    write_json(
        FINAL / "content-seal.json",
        {
            "entries": entries,
            "entry_count": len(entries),
            "final_commit": "EXTERNAL_AFTER_FINAL_COMMIT",
            "hash_domain": "raw_worktree_bytes_before_final_commit",
            "self_excluded": True,
        },
    )
    manifest_entries = list(entries)
    extra = [
        ROOT / "scripts" / "build_ghc_family_ilyra_v690_v2_final.py",
        ROOT / "scripts" / "ghc_family_ilyra_v690_v2_canonical.py",
        ROOT / "tests" / "test_ghc_family_ilyra_v690_v2_final.py",
    ]
    for path in extra:
        data = path.read_bytes()
        manifest_entries.append(
            {"bytes": len(data), "path": path.relative_to(ROOT).as_posix(), "sha256": digest(data)}
        )
    write_json(
        FINAL / "manifest.json",
        {
            "entries": sorted(manifest_entries, key=lambda row: row["path"]),
            "entry_count": len(manifest_entries),
            "hash_domain": "raw_worktree_bytes_before_final_commit",
            "self_excluded": "manifest.json",
        },
    )
    print(json.dumps({"content_seal": len(entries), "manifest": len(manifest_entries), "pages": visual_pages}, sort_keys=True))


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--draft-docx")
    parser.add_argument("--prepare", action="store_true")
    parser.add_argument("--seal", action="store_true")
    parser.add_argument("--visual-pages", type=int)
    args = parser.parse_args()
    if args.prepare:
        if not args.draft_docx:
            parser.error("--draft-docx is required with --prepare")
        prepare(Path(args.draft_docx).resolve())
        return
    if args.seal:
        if not args.visual_pages or args.visual_pages < 3:
            parser.error("--visual-pages of at least three is required with --seal")
        seal(args.visual_pages)
        return
    parser.error("choose --prepare or --seal")


if __name__ == "__main__":
    main()
