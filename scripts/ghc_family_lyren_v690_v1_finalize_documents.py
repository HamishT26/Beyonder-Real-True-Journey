"""Apply render-recovery accounting and seal Lyren's final documents."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
FINAL = ROOT / "docs/lyren-moss/v690-v1/final"
BOUNDARY = (
    "Bounded same-owner document rendering and visual review only; no affected-user accessibility "
    "evaluation, external audit, independent reproduction, production authority, or broader claim."
)


def raw_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text, flags=re.UNICODE))


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )


def replace_text(path: Path, replacements: dict[str, str]) -> None:
    text = path.read_text(encoding="utf-8")
    for old, new in replacements.items():
        if old not in text:
            if new in text:
                continue
            raise RuntimeError(f"replacement source absent in {path}: {old}")
        text = text.replace(old, new)
    path.write_text(text, encoding="utf-8", newline="\n")


def replace_docx_text(path: Path, replacements: dict[str, str]) -> None:
    document = Document(path)
    containers = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                containers.extend(cell.paragraphs)
    seen = {old: 0 for old in replacements}
    for paragraph in containers:
        for old, new in replacements.items():
            if old in paragraph.text:
                for run in paragraph.runs:
                    if old in run.text:
                        run.text = run.text.replace(old, new)
                        seen[old] += 1
                        break
                else:
                    paragraph.text = paragraph.text.replace(old, new)
                    seen[old] += 1
    missing = [old for old, count in seen.items() if count == 0]
    if missing:
        raise RuntimeError(f"DOCX replacement sources absent: {missing}")
    document.save(path)


def rebuild_baton() -> dict[str, object]:
    baton_root = FINAL / "baton"
    module_paths = sorted(baton_root.glob("*.md"))
    module_records = []
    combined_parts = [
        "# LYREN MOSS v690-v1 -> ILYRA FEN v690-v2 ACTIVATION CANDIDATE",
        "",
        "Committed delivery state: `PREPARED_NOT_SENT`.",
        "",
        "Read this complete file through the literal EOF marker before any mutation.",
        "",
    ]
    for path in module_paths:
        text = path.read_text(encoding="utf-8")
        module_records.append(
            {
                "path": path.relative_to(ROOT).as_posix(),
                "words": word_count(text),
                "bytes": path.stat().st_size,
                "sha256": raw_sha256(path),
            }
        )
        combined_parts.extend([text.rstrip(), ""])
    combined_parts.append("EOF LYREN MOSS v690-v1 BATON.")
    combined = "\n".join(combined_parts) + "\n"
    handoff = FINAL / "hand-off-baton.md"
    handoff.write_text(combined, encoding="utf-8", newline="\n")
    record = {
        "schema": "ghc.family.modular-baton-manifest.v1",
        "owner": "Lyren Moss",
        "phase": "v690-v1",
        "recipient": "Ilyra Fen",
        "recipient_phase": "v690-v2",
        "delivery_state": "PREPARED_NOT_SENT",
        "module_count": len(module_records),
        "modules": module_records,
        "combined": {
            "path": "docs/lyren-moss/v690-v1/final/hand-off-baton.md",
            "words": word_count(combined),
            "bytes": handoff.stat().st_size,
            "sha256": raw_sha256(handoff),
            "literal_last_line": "EOF LYREN MOSS v690-v1 BATON.",
        },
        "range": {
            "minimum_words": 10_000,
            "maximum_words": 100_000,
            "in_range": 10_000 <= word_count(combined) <= 100_000,
        },
    }
    write_json(FINAL / "baton-manifest.json", record)
    return record


def prepare() -> None:
    recovery_path = FINAL / "render-recovery.json"
    if recovery_path.exists():
        raise SystemExit("render recovery already prepared; refusing replay")
    replacements = {
        "five final-authoring and render-path failures": "six final-authoring, render-path, and correction failures",
        "two final-authoring failures": "six final-authoring, render-path, and correction failures",
        "972 effective negatives, 81 methods, and 2,477 direct witnesses: 683 failed and 1,794 passing": "973 effective negatives, 81 methods, and 2,479 direct witnesses: 684 failed and 1,795 passing",
        "969 effective negatives, 81 methods, and 2,471 direct witnesses: 680 failed and 1,791 passing": "973 effective negatives, 81 methods, and 2,479 direct witnesses: 684 failed and 1,795 passing",
    }
    replace_text(FINAL / "overview.md", replacements)
    replace_text(FINAL / "overview.html", replacements)
    replace_text(
        FINAL / "baton/05-retained-failures-and-methods.md",
        {
            "Two final-authoring wrapper failures add two negatives, one method, and four witnesses. The prepared terminal lineage is therefore 969 effective negatives, 81 methods, and 2,471 direct witnesses, of which 680 failed and 1,791 passed.": "Six final-authoring, render-path, and correction failures add six negatives, one method, and twelve witnesses. The prepared terminal lineage is therefore 973 effective negatives, 81 methods, and 2,479 direct witnesses, of which 684 failed and 1,795 passed."
        },
    )
    replace_text(
        FINAL / "baton/10-reports-and-accessibility.md",
        {
            "Rendering and visual inspection test only the generated artifact's local layout.": "Rendering and visual inspection test only the generated artifact's local layout. Three LibreOffice-dependent conversion attempts failed and remain retained; Microsoft Word exported a five-page PDF, and direct Poppler rasterization supplied the recovery path."
        },
    )
    replace_docx_text(
        FINAL / "overview.docx",
        {
            "two final-authoring failures": "six final-authoring, render-path, and correction failures",
            "969 effective negatives, 81 methods, and 2,471 direct witnesses: 680 failed and 1,791 passing": "973 effective negatives, 81 methods, and 2,479 direct witnesses: 684 failed and 1,795 passing",
            "969 negatives": "973 negatives",
            "2,471 witnesses": "2,479 witnesses",
            "680 failed": "684 failed",
            "1,791 passing": "1,795 passing",
        },
    )

    operational_path = FINAL / "operational-overlay.json"
    operational = load(operational_path)
    new_failures = [
        {
            "negative_id": "LM6901-FINAL-DOCX-RENDER-NO-LIBREOFFICE",
            "failure": "The prescribed DOCX renderer failed because LibreOffice soffice.exe was absent from PATH.",
            "recovery": "Use the installed Microsoft Word COM interface to export a five-page retained-preview PDF.",
            "original_success_credit": 0,
            "retained": True,
        },
        {
            "negative_id": "LM6901-FINAL-PDF-RASTER-AUTO-DPI",
            "failure": "The first PDF raster wrapper attempted an automatic DPI probe that still invoked the absent LibreOffice converter.",
            "recovery": "Attempt the renderer's explicit-DPI mode while retaining this failed wrapper at zero credit.",
            "original_success_credit": 0,
            "retained": True,
        },
        {
            "negative_id": "LM6901-FINAL-PDF-RASTER-EXPLICIT-DPI",
            "failure": "The explicit-DPI renderer path also invoked the absent LibreOffice converter and failed.",
            "recovery": "Stop retrying that renderer and use the bundled Poppler pdftoppm executable directly; five pages rasterized successfully.",
            "original_success_credit": 0,
            "retained": True,
        },
        {
            "negative_id": "LM6901-FINAL-DOCUMENT-CORRECTION-PARTIAL",
            "failure": "The first document-correction pass updated Markdown and HTML, then stopped because module 05 used a different count sentence; DOCX and ledgers were still unchanged.",
            "recovery": "Preserve the already-correct report edits, add exact old-or-new guards, and resume against the module-specific sentence without discarding any artifact.",
            "original_success_credit": 0,
            "retained": True,
        },
    ]
    operational["retained_negatives"].extend(new_failures)
    for row in new_failures:
        operational["witnesses"].extend(
            [
                {"witness_id": row["negative_id"] + "-W-FAIL", "result": "fail", "retained_negative_ids": [row["negative_id"]]},
                {"witness_id": row["negative_id"] + "-W-RECOVERY", "result": "pass", "retained_negative_ids": [row["negative_id"]]},
            ]
        )
    operational["counts"] = {
        "retained_negatives": 6,
        "methods": 1,
        "witnesses": 12,
        "failed_witnesses": 6,
        "passing_witnesses": 6,
    }
    write_json(operational_path, operational)

    accounting_path = FINAL / "terminal-accounting.json"
    accounting = load(accounting_path)
    accounting["layers"].append(
        {
            "layer": "Lyren render-path recovery within final authoring method",
            "effective_negatives": 4,
            "methods": 0,
            "direct_witnesses": 8,
            "failed_witnesses": 4,
            "passing_witnesses": 4,
        }
    )
    accounting["effective_total"] = {
        "effective_negatives": 973,
        "methods": 81,
        "direct_witnesses": 2479,
        "failed_witnesses": 684,
        "passing_witnesses": 1795,
    }
    write_json(accounting_path, accounting)

    write_json(
        recovery_path,
        {
            "schema": "ghc.family.document-render-recovery.v1",
            "owner": "Lyren Moss",
            "phase": "v690-v1-final",
            "retained_failures": [row["negative_id"] for row in new_failures],
            "retained_preview_pdf": "docs/lyren-moss/v690-v1/final/rendered/overview.pdf",
            "retained_preview_pages": 5,
            "correction": "Update the DOCX and accounting, then export a separately named corrected PDF and rasterize it directly with Poppler.",
            "corrected_pdf": "docs/lyren-moss/v690-v1/final/rendered/corrected/overview-final.pdf",
            "corrected_pages": "docs/lyren-moss/v690-v1/final/rendered/corrected/pages",
            "state": "CORRECTED_DOCX_PENDING_RENDER",
            "boundary": BOUNDARY,
        },
    )
    document_build_path = FINAL / "document-build.json"
    document_build = load(document_build_path)
    document_build["render_state"] = "CORRECTED_DOCX_PENDING_RENDER"
    document_build["retained_failed_render_attempts"] = 3
    document_build["retained_preview_pages"] = 5
    write_json(document_build_path, document_build)
    baton = rebuild_baton()
    print(
        json.dumps(
            {
                "state": "CORRECTED_DOCX_PENDING_RENDER",
                "baton_words": baton["combined"]["words"],
                "retained_failures": 6,
                "effective_negatives": 973,
            },
            sort_keys=True,
        )
    )


def seal() -> None:
    manifest_path = FINAL / "manifest.json"
    if manifest_path.exists():
        raise SystemExit("final manifest already exists; refusing replay")
    corrected = FINAL / "rendered/corrected"
    pdf = corrected / "overview-final.pdf"
    pages = sorted((corrected / "pages").glob("page-*.png"))
    if not pdf.is_file() or len(pages) != 5:
        raise RuntimeError("corrected PDF or five page images are unavailable")
    document_build_path = FINAL / "document-build.json"
    document_build = load(document_build_path)
    document_build["render_state"] = "VALIDATED_WITH_RETAINED_FAILURES"
    document_build["canonical_render"] = {
        "pdf": {"path": pdf.relative_to(ROOT).as_posix(), "bytes": pdf.stat().st_size, "sha256": raw_sha256(pdf), "pages": 5},
        "page_images": [{"path": path.relative_to(ROOT).as_posix(), "bytes": path.stat().st_size, "sha256": raw_sha256(path)} for path in pages],
        "converter": "Microsoft Word COM ExportAsFixedFormat",
        "rasterizer": "bundled Poppler pdftoppm",
    }
    document_build["visual_review"] = "all five corrected pages inspected"
    document_build["manual_affected_user_accessibility_review"] = "reserved"
    write_json(document_build_path, document_build)

    write_json(
        FINAL / "visual-review.json",
        {
            "schema": "ghc.family.visual-review.v1",
            "artifact": pdf.relative_to(ROOT).as_posix(),
            "pages": 5,
            "pages_inspected": [1, 2, 3, 4, 5],
            "clipping": False,
            "overlap": False,
            "unreadable_tables": False,
            "broken_heading_hierarchy_observed": False,
            "blank_required_content": False,
            "result": "pass",
            "same_owner_only": True,
            "affected_user_review": "reserved",
            "complete_accessibility_claimed": False,
            "boundary": BOUNDARY,
        },
    )
    recovery_path = FINAL / "render-recovery.json"
    recovery = load(recovery_path)
    recovery["state"] = "CORRECTED_PDF_AND_PAGES_VALIDATED"
    recovery["corrected_pdf_sha256"] = raw_sha256(pdf)
    recovery["corrected_page_sha256"] = [raw_sha256(path) for path in pages]
    write_json(recovery_path, recovery)

    content_seal_path = FINAL / "content-seal.json"
    content_entries = []
    for path in sorted(FINAL.rglob("*")):
        if path.is_file() and path.name not in {"content-seal.json", "manifest.json", "allowlist.json"}:
            content_entries.append(
                {
                    "path": path.relative_to(ROOT).as_posix(),
                    "bytes": path.stat().st_size,
                    "sha256": raw_sha256(path),
                }
            )
    write_json(
        content_seal_path,
        {
            "schema": "ghc.family.content-seal.v1",
            "hash_domain": "raw_worktree_bytes_before_final_commit",
            "self_excluded": True,
            "entries": content_entries,
            "entry_count": len(content_entries),
            "final_commit": "EXTERNAL_AFTER_FINAL_COMMIT",
        },
    )
    allowlist_path = FINAL / "allowlist.json"
    all_files = sorted(
        path.relative_to(ROOT).as_posix()
        for path in FINAL.rglob("*")
        if path.is_file()
    )
    all_files.extend(
        [
            allowlist_path.relative_to(ROOT).as_posix(),
            manifest_path.relative_to(ROOT).as_posix(),
        ]
    )
    write_json(
        allowlist_path,
        {
            "owner": "Lyren Moss",
            "phase": "final",
            "allowed_paths": sorted(set(all_files)),
            "additive_owner_scope": True,
        },
    )
    entries = []
    for path in sorted(file for file in FINAL.rglob("*") if file.is_file() and file != manifest_path):
        entries.append(
            {
                "path": path.relative_to(ROOT).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": raw_sha256(path),
                "hash_domain": "raw_file_bytes",
            }
        )
    write_json(
        manifest_path,
        {
            "schema": "ghc.family.owner-manifest.v1",
            "self_excluded": "manifest.json",
            "entry_count": len(entries),
            "entries": entries,
        },
    )
    print(
        json.dumps(
            {
                "state": "FINAL_DOCUMENTS_SEALED",
                "pages": len(pages),
                "manifest_entries": len(entries),
                "content_seal_entries": len(content_entries),
            },
            sort_keys=True,
        )
    )


def late() -> None:
    late_path = FINAL / "late-validation-overlay.json"
    if not late_path.is_file():
        raise RuntimeError("late validation overlay is unavailable")
    accounting_path = FINAL / "terminal-accounting.json"
    accounting = load(accounting_path)
    accounting["successor_visible_after_late_overlay"] = load(late_path)["successor_visible_accounting"]
    write_json(accounting_path, accounting)
    completion_path = FINAL / "completion-ledger.json"
    completion = load(completion_path)
    completion["late_validation_overlay"] = "docs/lyren-moss/v690-v1/final/late-validation-overlay.json"
    write_json(completion_path, completion)
    route_path = FINAL / "route-candidate.json"
    route = load(route_path)
    route["successor_visible_accounting"] = load(late_path)["successor_visible_accounting"]
    write_json(route_path, route)
    module_13 = FINAL / "baton/13-terminal-gates-and-recovery.md"
    marker = "## Late validation overlay"
    module_text = module_13.read_text(encoding="utf-8")
    if marker not in module_text:
        module_text += (
            "\n## Late validation overlay\n\n"
            "After the five-page document seal, the first Ruff correction pass for the final test and canonical "
            "validator fixed six findings but stopped on two remaining structural findings. That attempt retains zero "
            "success credit. The exact two-file recovery passed. Keep this layer separate from the rendered report's "
            "prepared accounting: the successor-visible overlay is 977 effective negatives, 82 methods, and 2,487 "
            "direct witnesses, with 688 failed and 1,799 passing. The terminal verdict remains "
            "`NOT_READY_FOR_STAGE_20`.\n"
        )
    else:
        module_text = re.sub(
            r"the successor-visible overlay is [0-9,]+ effective negatives, [0-9,]+ methods, and [0-9,]+ direct witnesses, with [0-9,]+ failed and [0-9,]+ passing",
            "the successor-visible overlay is 977 effective negatives, 82 methods, and 2,487 direct witnesses, with 688 failed and 1,799 passing",
            module_text,
        )
    module_13.write_text(module_text, encoding="utf-8", newline="\n")
    compact_path = FINAL / "compact-activation.md"
    compact = compact_path.read_text(encoding="utf-8")
    overlay_sentence = (
        "Successor-visible late accounting is 977 effective negatives, 82 methods, and 2,487 direct witnesses "
        "(688 failed / 1,799 passing); the five-page report preserves its earlier prepared snapshot and the late "
        "overlay remains separate.\n"
    )
    compact = re.sub(
        r"Successor-visible late accounting is [0-9,]+ effective negatives, [0-9,]+ methods, and [0-9,]+ direct witnesses \([0-9,]+ failed / [0-9,]+ passing\); the five-page report preserves its earlier prepared snapshot and the late overlay remains separate\.\n?",
        "",
        compact,
    ).rstrip()
    compact = compact + "\n\n" + overlay_sentence
    compact_path.write_text(compact, encoding="utf-8", newline="\n")
    baton = rebuild_baton()

    content_seal_path = FINAL / "content-seal.json"
    manifest_path = FINAL / "manifest.json"
    allowlist_path = FINAL / "allowlist.json"
    content_entries = []
    for path in sorted(FINAL.rglob("*")):
        if path.is_file() and path.name not in {"content-seal.json", "manifest.json", "allowlist.json"}:
            content_entries.append(
                {
                    "path": path.relative_to(ROOT).as_posix(),
                    "bytes": path.stat().st_size,
                    "sha256": raw_sha256(path),
                }
            )
    write_json(
        content_seal_path,
        {
            "schema": "ghc.family.content-seal.v1",
            "hash_domain": "raw_worktree_bytes_before_final_commit",
            "self_excluded": True,
            "entries": content_entries,
            "entry_count": len(content_entries),
            "final_commit": "EXTERNAL_AFTER_FINAL_COMMIT",
        },
    )
    actual_files = sorted(path.relative_to(ROOT).as_posix() for path in FINAL.rglob("*") if path.is_file())
    write_json(
        allowlist_path,
        {
            "owner": "Lyren Moss",
            "phase": "final",
            "allowed_paths": actual_files,
            "additive_owner_scope": True,
        },
    )
    entries = []
    for path in sorted(file for file in FINAL.rglob("*") if file.is_file() and file != manifest_path):
        entries.append(
            {
                "path": path.relative_to(ROOT).as_posix(),
                "bytes": path.stat().st_size,
                "sha256": raw_sha256(path),
                "hash_domain": "raw_file_bytes",
            }
        )
    write_json(
        manifest_path,
        {
            "schema": "ghc.family.owner-manifest.v1",
            "self_excluded": "manifest.json",
            "entry_count": len(entries),
            "entries": entries,
        },
    )
    print(
        json.dumps(
            {
                "state": "FINAL_LATE_OVERLAY_SEALED",
                "baton_words": baton["combined"]["words"],
                "manifest_entries": len(entries),
                "content_seal_entries": len(content_entries),
                "effective_negatives": 977,
            },
            sort_keys=True,
        )
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("mode", choices=["prepare", "seal", "late"])
    args = parser.parse_args()
    if args.mode == "prepare":
        prepare()
    elif args.mode == "seal":
        seal()
    else:
        late()


if __name__ == "__main__":
    main()
