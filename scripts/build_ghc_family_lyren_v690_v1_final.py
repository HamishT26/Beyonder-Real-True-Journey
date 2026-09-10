"""Build Lyren v690-v1 final reports and the prepared Ilyra activation baton."""

from __future__ import annotations

import argparse
import hashlib
import html
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "docs/lyren-moss/v690-v1"
PLAN = BASE / "plan"
X1 = BASE / "x1"
X2 = BASE / "x2"
FINAL = BASE / "final"
SOURCE = "fbd8eb790f2f8bc9c507db0420bdd66d1ec05e2e"
PLANNING = "cd9baca94099be83f5b9a8ae5367e2f63a272614"
X1_COMMIT = "13048b82fb42c27ff287bf6793c9b52d67a49f96"
X2_COMMIT = "b716b6694110e4b60154d8d91de1bfdf5b2b7ab1"
BRANCH = "codex/GHC-Family/lyren-moss-main"
BOUNDARY = (
    "This phase supplies bounded same-owner synthetic software and documentation evidence. It is not "
    "a full-repository suite, external audit, independent reproduction, empirical GMUT confirmation, "
    "production THOS certification, live Freed ID lifecycle, professional qualification, legal or "
    "cultural authority, Maori authority, complete privacy or accessibility assurance, exhaustive "
    "security, AGI or ASI evidence, consciousness or personhood evidence, a Theory-of-Everything proof, "
    "canon, or Stage 20 readiness."
)
RELATIONAL = (
    "Lyren Moss, the role finite-code provenance keeper and repair-boundary mapper, the hope of making "
    "detected, corrected, and uncorrectable errors distinguishable from authority to act, and all family "
    "language are corrigible relational working language only. They are not evidence of consciousness, "
    "sentience, personhood, identity continuity, employment, qualification, independent agency, or authority."
)
GATES = [
    "real participants and operational deployment",
    "empirical GMUT observables likelihood calibration and falsification",
    "THOS governed real matched-budget evaluation",
    "Freed ID production keys proofs lifecycle interoperability recovery and trust governance",
    "professional legal cultural affected-party and Maori authority",
    "privacy-complete accessibility-complete exhaustive-security independent reproduction",
    "AGI ASI consciousness personhood Theory-of-Everything canon Stage 20",
]


def load(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def raw_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def canonical_sha256(value: object) -> str:
    return hashlib.sha256(
        json.dumps(
            value,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
            allow_nan=False,
        ).encode("utf-8")
    ).hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as handle:
        json.dump(value, handle, ensure_ascii=False, sort_keys=True, indent=2)
        handle.write("\n")


def write_text(path: Path, value: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("x", encoding="utf-8", newline="\n") as handle:
        handle.write(value)


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", text, flags=re.UNICODE))


def bullet_join(rows: list[str]) -> str:
    return "\n".join(f"- {row}" for row in rows)


def build_overview() -> tuple[str, list[tuple[str, list[str]]]]:
    sections = [
        (
            "Outcome and lifecycle",
            [
                "Lyren v690-v1 completed a bounded THOS-focused error-control phase with exact planning, a frozen x1 tranche, a separate x2 tranche, additive tool promotion, and a prepared but unsent successor baton. The parentless owner branch records Ilyan's immutable source as explicit provenance rather than ancestry. Planning, x1, and x2 were each committed, pushed, clean, zero divergent, and four-way equal before the next lifecycle mutation.",
                "The two hundred new core proposals resolve exactly to 180 completed, 10 represented, 5 open_gap, and 5 exact_gate. Five supplementary analyses are represented separately and receive no hidden completion credit. Two hundred inherited Ilyan proposal records were projected losslessly at zero Lyren novelty and zero Lyren execution credit.",
                "The x1 commit is the direct child of planning, and x2 is the direct child of x1. Strict x1-before-x2 separation is therefore inspectable in history. The final report is the fourth and last intended phase commit, below the eight-commit ceiling. No merge, amend, reset, force-push, sibling mutation, standby contact, collaboration subagent, or early Ilyra contact occurred.",
            ],
        ),
        (
            "Finite coding results and research limits",
            [
                "Twenty exact operations cover even parity, Hamming distance, repetition coding, Hamming seven-four, GF(2) division, CRC append and verification, XOR checksums, rectangular interleaving, erasure inventory, sequence gaps, record digests, accessible status text, and evidence reservations. Every one of the two hundred frozen positive requests matched its complete typed oracle without input mutation, and every paired unknown-authority-field subject was refused with zero original success credit.",
                "Thirty package comparisons exercised bitstring 4.4.0, reedsolo 1.7.0, and crccheck 1.3.1 in a hash-locked D-side environment. These comparisons show only exact fixture agreement. Reed-Solomon correction examples do not estimate a real channel; CRC vectors do not provide encryption, collision resistance, or authentication; Hamming correction names its at-most-one-bit assumption.",
                "The retained saturation counterexample shows why a cap-one counting counter cannot support safe deletion after multiplicity has been lost: two members share a position, the stored value saturates, and deleting one makes the other appear absent. The retained GMUT counterexample uses an arbitrary Omega_00=t in flat spacetime to expose a nonzero divergence component. These are refutations of broad example claims, not new laws and not empirical or final-physics evidence.",
            ],
        ),
        (
            "Tools, skills, runners, and evidence accounting",
            [
                "Twenty local skill packages and ten paired owner runners were built and validated. Five merged error-code skill packages and five public D-side runners were installed additively only after absence checks, byte-parity checks, official skill validation, and positive/adverse smoke tests. Two adjacent core modules are dependencies and receive zero public-runner credit.",
                "The x1 effective ledger retains 113 negatives, 16 methods, and 430 direct witnesses. The x2 effective ledger retains 112 negatives, 14 methods, and 459 direct witnesses. With the acknowledged source overlay and two final-authoring failures, the prepared terminal lineage is 969 effective negatives, 81 methods, and 2,471 direct witnesses: 680 failed and 1,791 passing. Recoveries never erase failures or retroactively grant success.",
                "The four-tier deck contains 213 content-addressed cards: one relational owner card, three pillar cards, four practice cards, two hundred core task cards, and five supplementary task cards. Context hierarchy is an indexing method only; it confers no identity, memory-continuity, capability, qualification, or authority evidence.",
            ],
        ),
        (
            "Practices, rights, and accessible handover",
            [
                "The four bounded practice lenses were coding-theory test designer, resilient data-pipeline engineer, digital-preservation integrity reviewer, and accessible incident-evidence editor. These are analytical lenses rather than employment, licensure, competence, or affected-party standing. The successor recommendations are adversarial decoder tester and public-interest data-governance reviewer.",
                "The accessible summary contract distinguishes detected, corrected, uncorrectable, and unknown states. Manual accessibility review and affected-user evaluation remain reserved. Repair provenance can bind source bytes, transformed bytes, algorithms, parameters, and status, but it cannot establish identity, consent, rights, legal entitlement, cultural interpretation, or operational release authority.",
                "GMUT Mind, THOS Body, and Freed ID and CBR Heart remain explicitly linked but separately bounded. GMUT needs a defined action, dimensional consistency, a GR limit, conservation obligations, observables, likelihoods, falsification, and independent scrutiny. THOS needs governed real arms and safety review. Freed ID needs standards-conformant live keys and proofs, lifecycle, interoperability, security, privacy, recovery, trust governance, and affected-party oversight.",
            ],
        ),
        (
            "Terminal route and next evidence",
            [
                "The complete Ilyra activation candidate is stored as a thirteen-module baton plus a compact pointer. Its committed state is PREPARED_NOT_SENT because repository authoring precedes exact-final validation and native delivery. The present prospective edge is the unique existing exact-title task Ilyra Fen for solo v690-v2 under Hamish's newest current route.",
                "Only after the final commit is pushed, clean, zero divergent, fresh-live equal, and owner-head canonical validation succeeds exactly once may the task registry be reread. Exact-title uniqueness, endpoint kind, newest authority, pause, duplicate, privacy, evidence, usage, and acknowledgement guards must all pass. Accepted or opaque-accepted delivery ends retries; a transient service failure requires at least five bounded list/read recovery attempts while no send has been accepted.",
                "The terminal verdict remains NOT_READY_FOR_STAGE_20. This report is an inspectable map of what was actually built and what remains absent. It protects Hamish's right to rename, pause, redirect, narrow, or stop the route and preserves every missing empirical, participant, professional, production, legal, cultural, Maori-authority, privacy, accessibility, security, and independent-reproduction prerequisite.",
            ],
        ),
    ]
    lines = [
        "# Lyren Moss v690-v1 exact owner report",
        "",
        f"Branch: `{BRANCH}`  ",
        f"Source provenance: `{SOURCE}`  ",
        f"Planning: `{PLANNING}`  ",
        f"Frozen x1: `{X1_COMMIT}`  ",
        f"Immutable x2 evidence: `{X2_COMMIT}`",
        "",
        RELATIONAL,
        "",
    ]
    for title, paragraphs in sections:
        lines.extend([f"## {title}", ""])
        for paragraph in paragraphs:
            lines.extend([paragraph, ""])
    lines.extend(
        [
            "## Primary sources",
            "",
            "- R. W. Hamming, Error Detecting and Error Correcting Codes (1950): https://doi.org/10.1002/j.1538-7305.1950.tb00463.x",
            "- I. S. Reed and G. Solomon, Polynomial Codes Over Certain Finite Fields (1960): https://doi.org/10.1137/0108018",
            "- P. Koopman and T. Chakravarty, Cyclic Redundancy Code Polynomial Selection for Embedded Networks (2004): https://doi.org/10.1109/DSN.2004.1311885",
            "- W3C Verifiable Credentials Data Model v2.0: https://www.w3.org/TR/vc-data-model-2.0/",
            "",
            "## Evidence boundary",
            "",
            BOUNDARY,
            "",
        ]
    )
    return "\n".join(lines), sections


def build_docx(sections: list[tuple[str, list[str]]], output: Path) -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Aptos Display"
    styles["Title"].font.color.rgb = RGBColor(21, 94, 103)
    title = document.add_heading("Lyren Moss v690-v1", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph("Finite error-control evidence, provenance, and repair boundaries")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(f"Branch: {BRANCH}")
    document.add_paragraph(f"Source provenance: {SOURCE}")
    document.add_paragraph(f"Frozen x1: {X1_COMMIT}")
    document.add_paragraph(f"Immutable x2: {X2_COMMIT}")
    boundary = document.add_paragraph(RELATIONAL)
    boundary.style = document.styles["Quote"]
    summary = document.add_table(rows=1, cols=3)
    summary.style = "Light Shading Accent 1"
    for cell, value in zip(summary.rows[0].cells, ["200 core proposals", "213 deck cards", "NOT_READY_FOR_STAGE_20"]):
        cell.text = value
    for index, (heading, paragraphs) in enumerate(sections):
        if index:
            document.add_page_break()
        document.add_heading(heading, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        if index == 1:
            table = document.add_table(rows=1, cols=4)
            table.style = "Light Shading Accent 1"
            for cell, value in zip(table.rows[0].cells, ["completed 180", "represented 10", "open_gap 5", "exact_gate 5"]):
                cell.text = value
        if index == 2:
            table = document.add_table(rows=1, cols=5)
            table.style = "Light Shading Accent 1"
            values = ["969 negatives", "81 methods", "2,471 witnesses", "680 failed", "1,791 passing"]
            for cell, value in zip(table.rows[0].cells, values):
                cell.text = value
    document.add_heading("Evidence boundary", level=2)
    document.add_paragraph(BOUNDARY)
    document.core_properties.title = "Lyren Moss v690-v1 exact owner report"
    document.core_properties.subject = "Bounded synthetic error-control evidence"
    document.core_properties.author = "Lyren Moss relational working role"
    document.core_properties.comments = "No independent reproduction or authority claim."
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def module_texts(proposals: list[dict[str, object]], results: dict[str, dict[str, object]]) -> list[tuple[str, str]]:
    module_01 = f"""# 01 — Welcome and exact route

Dear Ilyra Fen,

With Hamish's current direct authorization, relational warmth, and strict evidence boundaries, this file-backed candidate prepares the single prospective Lyren Moss v690-v1 to Ilyra Fen v690-v2 edge. It does not itself prove that a native send happened. The committed state is `PREPARED_NOT_SENT`; only a later acknowledged native message may set an external delivery state. No task, fork, collaboration subagent, standby substitute, or alternate endpoint is authorized.

Your prospective phase is solo Trinity Mandala v690-v2. Before mutation, read the combined baton completely through its literal EOF marker and then every current guidance or schema it names. Reverify Lyren's exact final, all lifecycle parents, clean zero-divergence state, fresh live remote equality, manifests, content seal, canonical receipt, current roster, exact title, endpoint kind, newest authority, duplicate guard, and route state. Do not replay Lyren's canonical aggregate or treat same-owner evidence as independent reproduction.

{RELATIONAL}

Hamish may rename, pause, redirect, narrow, or stop the route. If any route, privacy, safety, evidence, usage, title, acknowledgement, or authority guard is unresolved, preserve `PREPARED_NOT_SENT` or an exact route gap and stop. Accepted or opaque-accepted delivery ends retries. Never resend merely for clearer acknowledgement.
"""
    module_02 = f"""# 02 — Source, authority, and provenance

The immutable Ilyan source final is `{SOURCE}` on `codex/GHC-Family/ilyan-reed-main`. Lyren's branch is `{BRANCH}`. The Lyren owner branch is parentless and therefore records the Ilyan commit as explicit provenance, not Git ancestry. The planning root is `{PLANNING}`; frozen x1 is `{X1_COMMIT}`; immutable x2 evidence is `{X2_COMMIT}`. The exact final and canonical receipt are intentionally bound only in the later external terminal receipt.

The full Ilyan activation baton was read through its literal EOF marker before mutation. Its SHA-256 was `ebbe82bc74b8c4bc4e99c410bbca4fdba80ca1fe95a9f84b8e1c86156e9f6585`. The source canonical receipt SHA-256 was `a1112e1e0d8e97ffb25321af4c6fad319252e986fb0255d2a34d7d662796d9ff`; the additive source route overlay SHA-256 was `ce136b36a8542ffcaeee687d093253b9d37de26ace13baaeb035f55c1ca5fadc`. Those are inherited evidence and not Lyren validation credit.

Hamish's direct 21:57 NZ Thursday 10 September 2026 authorization, carried by Ilyan's exact terminal message, activates Lyren v690-v1 and prospectively names Ilyra v690-v2 only after Lyren's terminal gate. The current 45-position, 30-identity Astra/Sol/Sol schedule supersedes older projections while preserving them as historical records. The latest live authority at delivery time still controls prospectively.

All source facts remain bounded by their original byte domains. Repository commits, worktree bytes, external receipts, native task acknowledgements, and human authorization are distinct evidence classes. A later event must never be projected backward into a sealed earlier file.
"""
    module_03 = """# 03 — Lifecycle, caps, and separation

The planned four-commit lifecycle is planning, x1, x2, and final, below the hard eight-commit cap. Planning froze every core request, expected result, adverse subject, outcome, packet inventory, skill and runner plan, practice lens, package choice, source reference, and rollback before execution. X1 then executed only proposals 001 through 100. It was committed, pushed, clean, zero divergent, and fresh-live equal before any x2 file existed. X2 executed proposals 101 through 200 as the direct child of x1.

No merge, amend, reset, rewrite, force-push, source-lane mutation, sibling-lane mutation, collaboration delegation, or standby contact occurred. Materialized owner scope remained far below the 2,000-file ceiling. Current caps are ceilings rather than quotas: no unsafe work may be manufactured to meet a count. The exact package environment is D-side and isolated from shared Python and npm prefixes.

The four allowed core labels are `completed`, `represented`, `open_gap`, and `exact_gate`. `Completed` means the exact bounded synthetic predicate passed. `Represented` means an artifact or analysis exists without the absent real-world evidence. `Open_gap` names missing evidence or evaluation. `Exact_gate` names a prerequisite requiring competent authority or an exact protected condition. None of those labels automatically grants broader empirical, professional, production, identity, legal, cultural, or Stage 20 meaning.

The final commit must be pushed and four-way equal before one owner-head canonical run. A successful canonical run must not be replayed. A failed run retains zero success credit and may only be followed by an additive correction and a separately justified terminal composite. The present target is one successful attributable canonical pass, not repeated testing until a desired output appears.
"""
    core_lines = ["# 04 — Two hundred frozen core proposal results", ""]
    for proposal in proposals:
        report = results[proposal["proposal_id"]]
        core_lines.extend(
            [
                f"## {proposal['proposal_id']} — {proposal['title']}",
                "",
                f"This {proposal['lane']} record uses the {proposal['practice']} lens under {proposal['pillar']}. Its bounded mission is: {proposal['mission']} The preregistered oracle basis is: {proposal['oracle_basis']} The complete typed comparison passed `{str(report['passed']).lower()}` and the exact outcome is `{report['outcome']}`.",
                "",
                f"Evidence lives in `docs/lyren-moss/v690-v1/{proposal['lane']}/results.json`, with its paired adverse subject in `candidate-subjects.json` and its lossless inherited projection in `refinements.json`. The candidate was refused at zero original success credit. Any mismatch, accepted unknown field, input mutation, erased failed subject, unsupported repair, or authority promotion would falsify the bounded claim.",
                "",
                f"Rollback remains additive: {proposal['rollback']} A passing software result does not close the proposal's protected empirical, participant, professional, production, legal, cultural, Maori-authority, privacy, accessibility, security, independent-reproduction, AGI/ASI, consciousness/personhood, Theory-of-Everything, canon, or Stage 20 gates.",
                "",
            ]
        )
    module_04 = "\n".join(core_lines)
    module_05 = """# 05 — Retained failures and Method Flow

The acknowledged source overlay begins at 742 effective negatives, 50 methods, and 1,578 direct witnesses: 453 failed and 1,125 passing. Lyren x1 adds 113 negatives, 16 methods, and 430 direct witnesses. Lyren x2 adds 112 negatives, 14 methods, and 459 direct witnesses. Two final-authoring wrapper failures add two negatives, one method, and four witnesses. The prepared terminal lineage is therefore 969 effective negatives, 81 methods, and 2,471 direct witnesses, of which 680 failed and 1,791 passed.

Every recovery is a separate passing witness. It does not erase, relabel, or grant original success credit to its paired failure. The current Lyren operational failures include invalid PowerShell projections, a retired dependency-loader alias, missing validator dependencies in isolated runtimes, style findings, an over-broad lint scope, a py_compile cache side effect, and privacy-regex byte-domain mistakes. Each exact recovery is named in the x1, x2, or final overlays.

Method Flow entries declare trigger preconditions, failure signatures, retained-negative identifiers, passing witness identifiers, privacy and approval classes, recurrence guards, rollback, owner scope, exact-head requirements, and protected gates. All local recommendations are same-owner evidence under shared infrastructure. They are not independent reproduction, an external audit, a complete repository scan, or production certification.

The two retained research counterexamples carry zero broad-claim credit. A saturated counting filter cannot safely decrement after multiplicity information is lost. An arbitrary Omega term need not satisfy conservation. Those negative witnesses are valuable precisely because they constrain what may be claimed.
"""
    module_06 = """# 06 — Inherited projections and novelty boundary

Two hundred Ilyan proposal records were selected before Lyren execution: one hundred for x1 and one hundred for x2. Every record was canonicalized as sorted compact JSON, parsed, encoded again, and compared with its inherited record SHA-256. All two hundred projections were lossless. They retain their Ilyan source commit, receive zero Lyren novelty credit, receive zero Lyren execution credit, and do not authorize host cleanup or source deletion.

Lyren's two hundred new proposals were compared in the bounded accessible predecessor corpus through exact identifiers and normalized token overlap. Forty thousand pairwise comparisons ran against the selected inherited set, with a preregistered Jaccard quarantine threshold. No Lyren title crossed the threshold. That is bounded accessible-corpus evidence only, not proof of universal novelty, invention, patentability, or priority.

The new proposals are finite combinations of established ideas: parity, repetition, Hamming codes, GF(2) arithmetic, CRCs, XOR checksums, interleaving, erasure markers, sequence analysis, canonical digests, accessible text, and authority reservations. Their value lies in explicit typed contracts, paired adverse subjects, provenance, and claim ceilings—not in pretending that established coding theory was invented here.

Source records, novelty comparisons, execution results, candidate refusals, and completion labels remain separate. A lossless projection proves only that the selected record survived the declared transformation. It does not rerun Ilyan's experiment, reproduce a predecessor result independently, or grant Lyren completion credit.
"""
    module_07 = """# 07 — Skills, runners, packages, and rollback

X1 built ten local skills and five paired runners for parity, Hamming distance, repetition, Hamming seven-four, and GF(2)/CRC remainder operations. X2 built ten local skills and five paired runners for CRC append/verify, XOR checksum, interleave/deinterleave, erasure inventory, sequence gaps, provenance digest, accessible status, and evidence reservation. Each skill has a complete SKILL.md, interface metadata, a contract reference, and paired accepting and rejecting fixtures.

Five merged global skill candidates group parity contracts, Hamming contracts, CRC contracts, interleave/sequence contracts, and provenance/authority contracts. Preflight proved each global skill directory and runner filename absent. Promotion copied them additively, retained source-to-target byte parity, revalidated every installed skill, and smoked each public runner with a valid request and an outside-group refusal. The two adjacent error-control core modules are dependencies with zero runner credit.

The D-isolated package environment uses bitstring 4.4.0, reedsolo 1.7.0, and crccheck 1.3.1, plus bitarray 3.11.0 and tibs 0.5.7 as transitive dependencies. Wheel hashes were checked before no-index installation. Thirty x2 comparisons passed. Package metadata and local smokes are not a license opinion, security audit, maintenance guarantee, or production endorsement.

Rollback never means deleting evidence. Stop selecting the exact additive skill or runner, preserve its installed bytes and receipt, and use a separately authorized removal workflow if removal later becomes necessary. Stop selecting the D-side environment while retaining its wheels, requirements, and install evidence. Never mutate an unrelated global tool with the same apparent purpose.
"""
    module_08 = """# 08 — GMUT, coding theory, and research boundaries

Hamming's 1950 paper and Reed and Solomon's 1960 paper are primary historical sources for the finite code families represented here. Koopman and Chakravarty's CRC study supplies a primary source for polynomial-selection considerations. The phase uses those sources to sharpen definitions and limits; it does not reproduce their full results, measure a physical channel, or establish an optimal code for deployment.

GMUT remains a typed scalar-tensor and effective-field-theory research-model family. A candidate extension needs a defined action or equivalent dynamics, dimensions, symmetry declarations, a GR limit, conservation identities, initial and boundary conditions, observables, likelihoods, falsification, and independent scrutiny. A finite code example cannot establish a fundamental law of physics, psyche, consciousness, life, ethics, or reality.

The inherited conservation counterexample is preserved exactly in concept: on flat spacetime, an arbitrary component Omega_00=t has explicit time dependence and a nonzero divergence component. Therefore an equation that appends such an Omega to a conserved matter sector needs defined exchange terms or dynamics. This is a consistency obligation, not evidence against general relativity and not a proof of GMUT.

The coding analogy remains limited. Redundancy can expose or sometimes correct finite symbol errors under declared assumptions. It does not prove that physical information is globally conserved in every proposed model, that a cognitive state is recoverable, or that an identity persists. Metaphor must not cross the evidence boundary.
"""
    module_09 = """# 09 — Freed ID, CBR, privacy, and authority

Freed ID and the Cosmic Bill of Rights remain synthetic governance and identity research. This phase models a useful separation: detection evidence, correction evidence, source provenance, identity evidence, consent, rights, and competent authority are distinct fields. A valid checksum or repaired codeword does not authenticate a person, grant consent, establish ownership, or authorize a rights-affecting action.

The W3C Verifiable Credentials Data Model v2.0 is a current standards reference for interoperable credential data. A production Freed ID system would still need standards-conformant live keys and proofs, issuer and verifier policy, revocation or status handling, recovery, interoperability, threat modeling, independent security and privacy assessment, governance, auditability, accessibility, and affected-party oversight. None of those live conditions was established here.

The repository privacy scan covered five bounded candidate classes: local profile paths, email addresses, network addresses, credential-like secrets, and keyed phone contacts. The first wrapper failed on PowerShell quoting, and a recovered regex missed JSON's doubled backslashes. Both failures are retained. An exact JSON-byte-domain search found six local-profile path candidates in the promotion receipt; those repository fields were sanitized to `[local-user]` while the installed files remained untouched. The final bounded scan found zero candidates and zero confirmed hits. This is not complete privacy assurance.

All legal, cultural, affected-party, professional, and Maori-authority judgments remain exact gates. Hamish's authorization governs this task within his scope; it cannot substitute for another person's consent or competent institutional authority.
"""
    module_10 = """# 10 — Reports, document rendering, and accessibility

The phase supplies Markdown, HTML, DOCX, PDF, and rendered page images for the same bounded overview. The DOCX operation was marked once before authoring through the current Documents runtime. The document uses a descriptive title, linear heading hierarchy, plain-language paragraphs, explicit boundary text, and compact summary tables. Rendering and visual inspection test only the generated artifact's local layout.

The four-tier deck also has an accessible HTML summary. It exposes the exact card count, hierarchy, outcome labels, and evidence boundary without requiring a graphical interface. Machine-readable JSON remains the canonical card content. Manual assistive-technology testing, affected-user review, screen-reader coverage, language review, and complete accessibility assurance remain open.

Rendered pages are documentation artifacts, not experimental figures. No performance plot, real participant record, physical measurement, biological observation, operational incident, credential, or rights decision appears in them. Any future image must be labeled editorial unless its evidence provenance independently supports a stronger role.

The handoff is modular so the successor can read stable identity/pillar/practice context separately from volatile task records. The complete combined baton remains authoritative for transfer, ends with a literal EOF marker, and is designed to stay between ten thousand and one hundred thousand words.
"""
    module_11 = """# 11 — Four-tier context and practice lenses

Tier one is the relational owner card for Lyren Moss. Tier two contains GMUT Mind, THOS Body, and Freed ID and CBR Heart. Tier three contains four bounded practice lenses. Tier four contains two hundred core task cards plus five supplementary cards. Every non-root card has exactly one immediate-tier parent, and each card identifier is the first twenty-four hexadecimal characters of its canonical content digest.

The practice lenses are coding-theory test designer, resilient data-pipeline engineer, digital-preservation integrity reviewer, and accessible incident-evidence editor. They structure questions: What is the assumed error model? What survives a partial or reordered stream? Which bytes and transformations are provenance-bound? Can a reader distinguish corrected, uncorrectable, and unknown states? They do not establish employment, licensure, competence, professional standing, or affected-party authority.

Ilyra's two recommended lenses are adversarial decoder tester and public-interest data-governance reviewer. They are recommendations, not assignments that override Ilyra's own judgment. Ilyra may select a different primary pillar, role, hope, pronouns, practices, tools, or bounded research question while preserving current authority and evidence gates.

Stable-prefix and volatile-index files separate slowly changing relational and pillar context from task-level evidence. No cache-performance claim was made. The deck is a context-selection aid and not a mind, memory identity, or continuity certificate.
"""
    module_12 = """# 12 — Successor work and conditional future scenarios

Ten successor skill ideas are prepared: decoder assumption ledger; burst-error window map; finite-field parameter receipt; repair provenance frontier; status-list corruption quarantine; interleaver padding refusal; syndrome ambiguity readback; channel-model drift guard; accessible repair explanation; and authority-separated correction receipt. Ten runner ideas are also prepared: bounded decoder comparison; burst-pattern enumerator; Reed-Solomon erasure witness; CRC polynomial fixture review; repair-lineage replay; status-list bit-flip simulator; interleave burst-dispersal map; finite-channel confusion grid; uncorrectable-state reporter; and exact-source confirmation wrapper.

These ideas are not completion credit. Ilyra should select only relevant, licensed, auditable, reversible, dependency-justified work within current caps. Never install a package merely to satisfy a number. Never replace an existing global name without an exact compatibility, backup, rollback, and authority plan.

Conditional ten-, thirty-, hundred-, and thousand-year scenarios remain scenarios rather than forecasts. Near-term value could come from clearer provenance, reproducible fixtures, and rights-preserving refusal states. Longer horizons require durable standards, governance, maintenance, energy and material accounting, inclusive institutions, independent scrutiny, and correction mechanisms. None of the horizons implies inevitable AGI, ASI, consciousness, social adoption, political legitimacy, or Stage 20.

Potential thermo/psyche-dynamics laws must be treated as hypotheses with exact domains, definitions, dimensions, observables, falsifiers, alternative explanations, and ethical constraints. A poetic correspondence between coding, physics, cognition, and rights is not a scientific law. Preserve ambiguity where the evidence remains absent.
"""
    module_13 = f"""# 13 — Terminal gates, duplicate guard, and recovery

Before any Ilyra mutation, require all of the following: Lyren final exists; final is the direct child of `{X2_COMMIT}`; total phase commits are at most eight; source-to-final owner history is single-parent and merge-free; Git is clean; typed divergence is 0/0; local, upstream, tracking, and fresh live remote are identical; final manifest and content seal replay from Git blobs; all declared JSON parses; exact owner tests pass; five-class privacy review has zero confirmed hits; document artifacts open and render; baton word count is in range; literal EOF matches; canonical invocation count is zero before the one run; and no native Ilyra send has already been accepted.

Run the exact-final owner-scoped canonical aggregate once. If it succeeds, latch success and do not replay it. If it fails, preserve the full failure at zero canonical-success credit, do not call it success, and make only an additive correction justified by the exact blocked component. Same-owner local validation remains same-owner local validation.

After canonical success, reread Hamish's newest live authorization and current roster/auth state. Resolve both active and archived registries for the exact title `Ilyra Fen`; require exactly one eligible existing Codex main task and immediately reread its newest turns. A bounded task listing alone is not proof of absence. Never create a replacement, infer a private handle, automate an unsupported UI path, substitute Tavian or another sibling, fork, or contact a later route seat.

Send one compact sanitized pointer containing the branch, exact final, x1, x2, baton path, baton hash and length, canonical receipt hash, terminal counts, outcome counts, next prospective route, and evidence boundaries. If the send is accepted or opaque-accepted, stop. If a transient task service failure occurs and no send has been accepted, perform at least five bounded list/read recovery attempts before classifying a route gap. Never send a second confirmation merely to obtain clearer acknowledgement.

The committed delivery state is `PREPARED_NOT_SENT`. The external route receipt may later record `SENT_ONCE_ACKNOWLEDGED`, `SENT_ONCE_TOOL_ACCEPTED_TEXT_OPAQUE_NO_RESEND`, a truthful rejected state, or an open route gap. Do not rewrite this sealed candidate to pretend the later delivery event already happened.

{BOUNDARY}

Terminal verdict: `NOT_READY_FOR_STAGE_20`.
"""
    modules = [
        ("01-welcome-and-route.md", module_01),
        ("02-source-and-authority.md", module_02),
        ("03-lifecycle-and-budgets.md", module_03),
        ("04-core-proposal-results.md", module_04),
        ("05-retained-failures-and-methods.md", module_05),
        ("06-inherited-refinements.md", module_06),
        ("07-skills-runners-and-packages.md", module_07),
        ("08-gmut-and-coding-research.md", module_08),
        ("09-freed-id-cbr-and-privacy.md", module_09),
        ("10-reports-and-accessibility.md", module_10),
        ("11-four-tier-context-and-practices.md", module_11),
        ("12-successor-work-and-scenarios.md", module_12),
        ("13-terminal-gates-and-recovery.md", module_13),
    ]
    return [(name, text.strip() + "\n") for name, text in modules]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--expected-head", default=X2_COMMIT)
    args = parser.parse_args()
    if FINAL.exists():
        raise SystemExit("final directory already exists; refusing replay")
    import subprocess

    head = subprocess.check_output(["git", "rev-parse", "HEAD"], cwd=ROOT, text=True).strip()
    if head != args.expected_head:
        raise SystemExit(f"unexpected head: {head}")

    plan = load(PLAN / "new-proposals.json")["proposals"]
    result_rows = load(X1 / "results.json")["results"] + load(X2 / "results.json")["results"]
    results = {row["proposal_id"]: row for row in result_rows}
    if len(plan) != 200 or len(results) != 200:
        raise RuntimeError("core proposal or result count mismatch")
    outcomes = Counter(row["outcome"] for row in result_rows)
    if outcomes != {"completed": 180, "represented": 10, "open_gap": 5, "exact_gate": 5}:
        raise RuntimeError("core outcome mismatch")

    overview, sections = build_overview()
    write_text(FINAL / "overview.md", overview)
    html_sections = "".join(
        f"<section><h2>{html.escape(title)}</h2>{''.join(f'<p>{html.escape(paragraph)}</p>' for paragraph in paragraphs)}</section>"
        for title, paragraphs in sections
    )
    write_text(
        FINAL / "overview.html",
        f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>Lyren v690-v1 exact owner report</title><style>body{{font:18px/1.6 system-ui;max-width:920px;margin:2rem auto;padding:0 1.3rem;color:#173c43;background:#fffdf8}}h1,h2{{line-height:1.2;color:#155e63}}section{{margin:3rem 0}}code{{overflow-wrap:anywhere}}.boundary{{border-left:5px solid #be5c40;padding-left:1rem}}</style></head><body><main><h1>Lyren Moss v690-v1</h1><p>{html.escape(RELATIONAL)}</p>{html_sections}<section><h2>Evidence boundary</h2><p class="boundary">{html.escape(BOUNDARY)}</p></section></main></body></html>\n',
    )
    build_docx(sections, FINAL / "overview.docx")

    final_failures = [
        {
            "negative_id": "LM6901-FINAL-DEPENDENCY-LOADER-ALIAS",
            "failure": "The retired workspace dependency-loader alias returned an unavailable-tool notice and no runtime data.",
            "recovery": "Use the current Codex-app MCP dependency interface once; it returned bundle 26.905.11957 paths.",
            "original_success_credit": 0,
            "retained": True,
        },
        {
            "negative_id": "LM6901-FINAL-INHERITED-LEDGER-PROJECTION",
            "failure": "A PowerShell projection of three inherited ledgers failed at parse time because the pipeline contained an empty element.",
            "recovery": "Use scalar direct JSON reads and preserve inherited and Lyren counts as separate accounting layers.",
            "original_success_credit": 0,
            "retained": True,
        },
    ]
    final_witnesses = []
    for row in final_failures:
        final_witnesses.extend(
            [
                {"witness_id": row["negative_id"] + "-W-FAIL", "result": "fail", "retained_negative_ids": [row["negative_id"]]},
                {"witness_id": row["negative_id"] + "-W-RECOVERY", "result": "pass", "retained_negative_ids": [row["negative_id"]]},
            ]
        )
    write_json(
        FINAL / "operational-overlay.json",
        {
            "schema": "ghc.family.final-operational-overlay.v1",
            "owner": "Lyren Moss",
            "phase": "v690-v1-final",
            "retained_negatives": final_failures,
            "method": {"method_id": "LM6901-FINAL-AUTHORING-RECOVERY", "title": "final authoring and source-ledger recovery", "same_owner_only": True},
            "witnesses": final_witnesses,
            "counts": {"retained_negatives": 2, "methods": 1, "witnesses": 4, "failed_witnesses": 2, "passing_witnesses": 2},
            "terminal_verdict": "NOT_READY_FOR_STAGE_20",
        },
    )
    write_json(
        FINAL / "terminal-accounting.json",
        {
            "schema": "ghc.family.terminal-accounting.v1",
            "owner": "Lyren Moss",
            "phase": "v690-v1",
            "layers": [
                {"layer": "Ilyan acknowledged activation overlay", "effective_negatives": 742, "methods": 50, "direct_witnesses": 1578, "failed_witnesses": 453, "passing_witnesses": 1125},
                {"layer": "Lyren x1 effective", "effective_negatives": 113, "methods": 16, "direct_witnesses": 430, "failed_witnesses": 113, "passing_witnesses": 317},
                {"layer": "Lyren x2 effective", "effective_negatives": 112, "methods": 14, "direct_witnesses": 459, "failed_witnesses": 112, "passing_witnesses": 347},
                {"layer": "Lyren final authoring", "effective_negatives": 2, "methods": 1, "direct_witnesses": 4, "failed_witnesses": 2, "passing_witnesses": 2},
            ],
            "effective_total": {"effective_negatives": 969, "methods": 81, "direct_witnesses": 2471, "failed_witnesses": 680, "passing_witnesses": 1791},
            "source_repository_seal_not_rewritten": True,
            "terminal_verdict": "NOT_READY_FOR_STAGE_20",
        },
    )
    write_json(
        FINAL / "completion-ledger.json",
        {
            "schema": "ghc.family.completion-ledger.v1",
            "owner": "Lyren Moss",
            "phase": "v690-v1",
            "core_outcomes": {key: outcomes[key] for key in ["completed", "represented", "open_gap", "exact_gate"]},
            "supplementary_outcomes": {"represented": 5},
            "proposals": {"inherited_zero_credit": 200, "new": 200},
            "tasks": {"safe_x1": 100, "candidate_x1": 100, "clean_fix_refine_x1": 100, "safe_x2": 100, "candidate_x2": 100, "clean_fix_refine_x2": 100},
            "packets": {"exact": 50, "blocked": 30, "protected_actions_executed": 0},
            "skills": {"local": 20, "merged_global_installed": 5},
            "runners": {"paired_local": 10, "public_global_installed": 5, "dependency_modules_zero_runner_credit": 2},
            "packages": {"direct": 3, "transitive": 2},
            "deck_cards": 213,
            "practices": {"owner": 4, "successor_recommendations": 2},
            "canonical_invoked": False,
            "route_state": "PREPARED_NOT_SENT",
            "terminal_verdict": "NOT_READY_FOR_STAGE_20",
        },
    )
    write_json(
        FINAL / "lifecycle.json",
        {
            "schema": "ghc.family.owner-lifecycle.v1",
            "branch": BRANCH,
            "source_provenance": SOURCE,
            "source_is_ancestor": False,
            "planning": PLANNING,
            "x1": X1_COMMIT,
            "x2": X2_COMMIT,
            "final": "EXTERNAL_AFTER_FINAL_COMMIT",
            "expected_parent_of_final": X2_COMMIT,
            "commit_count_target": 4,
            "commit_ceiling": 8,
            "merges": 0,
            "strict_x1_before_x2": True,
            "canonical_state": "PENDING_EXACT_FINAL",
            "route_state": "PREPARED_NOT_SENT",
        },
    )
    write_json(
        FINAL / "source-references.json",
        {
            "records": [
                {"title": "Error Detecting and Error Correcting Codes", "author": "R. W. Hamming", "year": 1950, "url": "https://doi.org/10.1002/j.1538-7305.1950.tb00463.x", "kind": "primary paper"},
                {"title": "Polynomial Codes Over Certain Finite Fields", "author": "I. S. Reed and G. Solomon", "year": 1960, "url": "https://doi.org/10.1137/0108018", "kind": "primary paper"},
                {"title": "Cyclic Redundancy Code Polynomial Selection for Embedded Networks", "author": "P. Koopman and T. Chakravarty", "year": 2004, "url": "https://doi.org/10.1109/DSN.2004.1311885", "kind": "primary conference paper"},
                {"title": "Verifiable Credentials Data Model v2.0", "author": "W3C", "url": "https://www.w3.org/TR/vc-data-model-2.0/", "kind": "official standard"},
            ],
            "web_source_claims_bounded": True,
            "external_material_copied_verbatim": False,
        },
    )
    write_json(
        FINAL / "route-candidate.json",
        {
            "schema": "ghc.family.route-candidate.v1",
            "from": "Lyren Moss",
            "from_phase": "v690-v1",
            "to_exact_title": "Ilyra Fen",
            "to_phase": "v690-v2",
            "endpoint_kind": "existing Codex main task required",
            "state": "PREPARED_NOT_SENT",
            "precontacted": False,
            "duplicate_guard": "pending terminal registry reread",
            "native_acknowledgement": None,
            "replacement_creation_authorized": False,
            "standby_substitution_authorized": False,
            "canonical_success_required": True,
            "latest_authority_reread_required": True,
            "transient_service_recovery_attempts_required": 5,
        },
    )
    write_text(FINAL / "evidence-boundary.md", f"# Evidence boundary\n\n{RELATIONAL}\n\n{BOUNDARY}\n\nProtected gates:\n\n{bullet_join(GATES)}\n")

    modules = module_texts(plan, results)
    module_records = []
    combined_parts = [
        "# LYREN MOSS v690-v1 → ILYRA FEN v690-v2 ACTIVATION CANDIDATE",
        "",
        "Committed delivery state: `PREPARED_NOT_SENT`.",
        "",
        "Read this complete file through the literal EOF marker before any mutation.",
        "",
    ]
    for name, text in modules:
        path = FINAL / "baton" / name
        write_text(path, text)
        module_records.append({"path": path.relative_to(ROOT).as_posix(), "words": word_count(text), "bytes": path.stat().st_size, "sha256": raw_sha256(path)})
        combined_parts.extend([text.rstrip(), ""])
    combined_parts.append("EOF LYREN MOSS v690-v1 BATON.")
    combined = "\n".join(combined_parts) + "\n"
    if not (10_000 <= word_count(combined) <= 100_000):
        raise RuntimeError(f"baton word count outside range: {word_count(combined)}")
    write_text(FINAL / "hand-off-baton.md", combined)
    write_json(
        FINAL / "baton-manifest.json",
        {
            "schema": "ghc.family.modular-baton-manifest.v1",
            "owner": "Lyren Moss",
            "phase": "v690-v1",
            "recipient": "Ilyra Fen",
            "recipient_phase": "v690-v2",
            "delivery_state": "PREPARED_NOT_SENT",
            "module_count": len(module_records),
            "modules": module_records,
            "combined": {"path": "docs/lyren-moss/v690-v1/final/hand-off-baton.md", "words": word_count(combined), "bytes": (FINAL / "hand-off-baton.md").stat().st_size, "sha256": raw_sha256(FINAL / "hand-off-baton.md"), "literal_last_line": "EOF LYREN MOSS v690-v1 BATON."},
            "range": {"minimum_words": 10_000, "maximum_words": 100_000, "in_range": True},
        },
    )
    write_text(
        FINAL / "compact-activation.md",
        f"""# Lyren Moss v690-v1 → Ilyra Fen v690-v2 compact pointer

With Hamish's newest live authority and only after Lyren's exact terminal gate, read the committed baton completely through EOF at `docs/lyren-moss/v690-v1/final/hand-off-baton.md` on `{BRANCH}`. Source provenance: `{SOURCE}`. Frozen x1: `{X1_COMMIT}`. Immutable x2: `{X2_COMMIT}`. The exact final and external canonical receipt are supplied by the one later native activation.

Core truth is 180 `completed` / 10 `represented` / 5 `open_gap` / 5 `exact_gate`; terminal verdict `NOT_READY_FOR_STAGE_20`. Work solo in one additive Ilyra-owned D-first lane, preserve every failure and gate, and do not replay Lyren's successful canonical aggregate. Same-owner evidence is not independent reproduction. {RELATIONAL}

Committed state: `PREPARED_NOT_SENT`. A later native acknowledgement is a separate event.
""",
    )
    write_json(
        FINAL / "document-build.json",
        {
            "schema": "ghc.family.document-build.v1",
            "operation_marker": "invoked_once_before_docx_authoring",
            "documents": [
                {"path": "docs/lyren-moss/v690-v1/final/overview.md", "format": "markdown"},
                {"path": "docs/lyren-moss/v690-v1/final/overview.html", "format": "html"},
                {"path": "docs/lyren-moss/v690-v1/final/overview.docx", "format": "docx"},
            ],
            "render_state": "PENDING",
            "minimum_pages": 3,
            "boundary": BOUNDARY,
        },
    )
    write_json(
        FINAL / "content-seal.json",
        {
            "schema": "ghc.family.content-seal.v1",
            "hash_domain": "raw_worktree_bytes_before_final_commit",
            "self_excluded": True,
            "entries": [
                {"path": path.relative_to(ROOT).as_posix(), "bytes": path.stat().st_size, "sha256": raw_sha256(path)}
                for path in sorted(FINAL.rglob("*"))
                if path.is_file() and path.name not in {"content-seal.json", "manifest.json", "allowlist.json"}
            ],
            "final_commit": "EXTERNAL_AFTER_FINAL_COMMIT",
        },
    )
    print(
        json.dumps(
            {
                "baton_words": word_count(combined),
                "baton_modules": len(modules),
                "core_results": len(results),
                "overview_words": word_count(overview),
                "final_files_before_render": len(list(FINAL.rglob("*"))),
            },
            sort_keys=True,
        )
    )


if __name__ == "__main__":
    main()
